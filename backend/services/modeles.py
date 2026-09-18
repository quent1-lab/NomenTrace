"""Modèles Excel à remplir par l'équipe : un par bloc fonctionnel ou par ensemble.

Le modèle reprend les composants existants avec leurs valeurs actuelles et laisse des lignes
vides pour les ajouts. Une feuille cachée l'identifie (type, code, date) pour que l'import le
reconnaisse quel que soit le nom du fichier.
"""

import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.worksheet.worksheet import Worksheet

from backend import db
from backend.erreurs import Introuvable
from backend.services import parametres
from backend.services.import_colonnes import (
    COLONNE_ENSEMBLE,
    COLONNE_QTE_ENSEMBLE,
    COLONNES,
    NOM_FEUILLE,
    NOM_FEUILLE_LISTES,
    NOM_FEUILLE_META,
    Colonne,
)

LIGNES_VIDES = 50
FOND_ENTETE = PatternFill("solid", fgColor="2A3B52")
FOND_ID = PatternFill("solid", fgColor="D9DEE5")
TEXTE_ID = Font(color="5F6B7A")
# Listes où une valeur nouvelle est admise : elle sera proposée à la création à l'import.
LISTES_OUVERTES: frozenset[str] = frozenset(
    {"fournisseur_nom", "ensemble_code", "mode_appro", "statut_appro", "statut_choix", "criticite"}
)


def _libelles(conn: sqlite3.Connection) -> dict[str, dict[str, str]]:
    """Libellés des listes paramétrables, par liste puis par code."""
    libelles: dict[str, dict[str, str]] = {}
    for v in db.fetch_all(conn, "SELECT liste, code, libelle FROM valeur_liste ORDER BY ordre"):
        libelles.setdefault(v["liste"], {})[v["code"]] = v["libelle"]
    return libelles


def _valeur_cellule(colonne: Colonne, composant: dict, libelles: dict) -> Any:
    valeur = composant.get(colonne.champ)
    if colonne.nature == "liste" and valeur is not None:
        return libelles.get(colonne.champ, {}).get(valeur, valeur)
    return valeur


def _ecrire(feuille: Worksheet, ligne: int, colonne: int, valeur: Any) -> None:
    cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
    if isinstance(valeur, str):
        cellule.data_type = "s"  # un texte commençant par « = » ne devient pas une formule


def _feuille_listes(classeur: Workbook, conn: sqlite3.Connection, libelles: dict) -> dict[str, str]:
    """Feuille cachée portant les valeurs des listes déroulantes ; renvoie les plages."""
    feuille = classeur.create_sheet(NOM_FEUILLE_LISTES)
    sources = {
        "bloc_code": [
            b["code"] for b in db.fetch_all(conn, "SELECT code FROM bloc ORDER BY ordre")
        ],
        "fournisseur_nom": [
            f["nom"]
            for f in db.fetch_all(
                conn, "SELECT nom FROM fournisseur WHERE archive = 0 ORDER BY nom COLLATE NOCASE"
            )
        ],
        "base_prix_releve": ["HT", "TTC"],
        "ensemble_code": [
            e["code"]
            for e in db.fetch_all(
                conn, "SELECT code FROM ensemble WHERE archive = 0 ORDER BY ordre"
            )
        ],
        **{liste: list(valeurs.values()) for liste, valeurs in libelles.items()},
    }
    plages = {}
    for index, (champ, valeurs) in enumerate(sources.items(), start=1):
        lettre = get_column_letter(index)
        feuille.cell(row=1, column=index, value=champ)
        for rang, valeur in enumerate(valeurs, start=2):
            _ecrire(feuille, rang, index, valeur)
        plages[champ] = f"{NOM_FEUILLE_LISTES}!${lettre}$2:${lettre}${max(2, len(valeurs) + 1)}"
    feuille.sheet_state = "hidden"
    return plages


def _entetes(feuille: Worksheet, colonnes: list[Colonne]) -> None:
    for index, colonne in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=1, column=index, value=colonne.entete)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = FOND_ENTETE
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
        feuille.column_dimensions[get_column_letter(index)].width = colonne.largeur
    feuille.cell(row=1, column=1).comment = Comment(
        "Colonne verrouillée : l'identifiant est attribué par Nomentrace. "
        "Pour un nouveau composant, remplir une ligne vide sans toucher à l'ID.",
        "Nomentrace",
    )
    feuille.freeze_panes = "C2"
    feuille.row_dimensions[1].height = 32


def _validations(feuille: Worksheet, colonnes: list[Colonne], plages: dict, derniere: int) -> None:
    for index, colonne in enumerate(colonnes, start=1):
        plage = plages.get(colonne.champ)
        if plage is None:
            continue
        validation = DataValidation(type="list", formula1=f"={plage}", allow_blank=True)
        validation.errorTitle = colonne.entete
        if colonne.champ in LISTES_OUVERTES:
            validation.errorStyle = "warning"
            validation.error = (
                "Valeur absente de la liste. Elle sera proposée à la création lors de l'import : "
                "« Oui » pour la garder, « Non » pour corriger."
            )
        else:
            validation.error = "Choisir une valeur de la liste."
        lettre = get_column_letter(index)
        validation.add(f"{lettre}2:{lettre}{derniere}")
        feuille.add_data_validation(validation)


def _verrouiller_id(feuille: Worksheet, nb_colonnes: int, derniere: int) -> None:
    """Colonne ID grisée et verrouillée ; toutes les autres cellules restent modifiables.

    La protection est sans mot de passe : elle évite une saisie par erreur, pas plus.
    """
    for ligne in feuille.iter_rows(min_row=2, max_row=derniere, max_col=nb_colonnes):
        for cellule in ligne[1:]:
            cellule.protection = Protection(locked=False)
        ligne[0].fill = FOND_ID
        ligne[0].font = TEXTE_ID
    feuille.protection = SheetProtection(
        sheet=True,
        formatCells=False,
        formatColumns=False,
        formatRows=False,
        sort=False,
        autoFilter=False,
        selectLockedCells=False,
        selectUnlockedCells=False,
    )


def _consignes(classeur: Workbook, titre: str, ensemble: bool) -> None:
    feuille = classeur.create_sheet("Consignes")
    lignes = [
        titre,
        "",
        "Une ligne par composant. Les lignes existantes portent leurs valeurs actuelles : "
        "modifier ce qui a changé, ne rien toucher sinon.",
        "Colonne ID (grisée) : verrouillée, c'est Nomentrace qui attribue les identifiants.",
        "Nouveau composant : utiliser une ligne vide, remplir au minimum Fonction, Désignation, "
        "Mode appro et Qté besoin" + (", ainsi que le Bloc." if ensemble else "."),
        "Fournisseur, Mode appro, statuts, criticité"
        + ("" if ensemble else ", Ensemble")
        + " : une valeur absente de la liste est acceptée ; sa création sera proposée à l'import.",
        "Réf fabricant : la renseigner dès qu'elle est connue, elle sert à repérer les doublons.",
        "PU relevé : le prix tel qu'affiché, en indiquant dans Base prix s'il est HT ou TTC.",
        "Les colonnes à valeurs fermées proposent une liste déroulante.",
        "Rien n'est enregistré directement : chaque proposition est relue avant d'être appliquée.",
    ]
    if not ensemble:
        lignes.insert(
            6,
            "Ensemble et Qté dans cet ensemble : facultatifs, pour indiquer où le composant est "
            "monté et en quelle quantité.",
        )
    if ensemble:
        lignes.insert(
            6,
            "Qté dans cet ensemble : la quantité réellement montée dans cette partie. "
            "0 ou vide retire le composant de l'ensemble.",
        )
    for rang, texte in enumerate(lignes, start=1):
        feuille.cell(row=rang, column=1, value=texte).font = Font(bold=rang == 1)
    feuille.column_dimensions["A"].width = 110


def _meta(classeur: Workbook, conn: sqlite3.Connection, type_modele: str, code: str) -> None:
    feuille = classeur.create_sheet(NOM_FEUILLE_META)
    valeurs = {
        "type": type_modele,
        "code": code,
        "genere_le": datetime.now().isoformat(timespec="seconds"),
        "projet": parametres.get_parametre(conn, "nom_projet") or "",
    }
    for rang, (cle, valeur) in enumerate(valeurs.items(), start=1):
        feuille.cell(row=rang, column=1, value=cle)
        _ecrire(feuille, rang, 2, valeur)
    feuille.sheet_state = "hidden"


def _construire(
    conn: sqlite3.Connection, lignes: list[dict], type_modele: str, code: str, titre: str
) -> Workbook:
    ensemble = type_modele == "ensemble"
    colonnes = list(COLONNES)
    if ensemble:
        colonnes.insert(4, COLONNE_QTE_ENSEMBLE)
    else:
        colonnes.extend((COLONNE_ENSEMBLE, COLONNE_QTE_ENSEMBLE))
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = NOM_FEUILLE
    libelles = _libelles(conn)
    _entetes(feuille, colonnes)
    for rang, composant in enumerate(lignes, start=2):
        for index, colonne in enumerate(colonnes, start=1):
            _ecrire(feuille, rang, index, _valeur_cellule(colonne, composant, libelles))
    derniere = len(lignes) + 1 + LIGNES_VIDES
    _verrouiller_id(feuille, len(colonnes), derniere)
    for index, colonne in enumerate(colonnes, start=1):
        lettre = get_column_letter(index)
        if colonne.nature == "montant":
            for (cellule,) in feuille[f"{lettre}2:{lettre}{derniere}"]:
                cellule.number_format = "0.00"
        if colonne.nature == "taux":
            for (cellule,) in feuille[f"{lettre}2:{lettre}{derniere}"]:
                cellule.number_format = "0%"
    _validations(feuille, colonnes, _feuille_listes(classeur, conn, libelles), derniere)
    _consignes(classeur, titre, ensemble)
    _meta(classeur, conn, type_modele, code)
    return classeur


def _enregistrer(classeur: Workbook, dossier: Path, type_modele: str, code: str) -> Path:
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = dossier / f"modele_{type_modele}_{code}_{datetime.now():%Y%m%d}.xlsx"
    classeur.save(chemin)
    return chemin


def generate_modele_bloc(conn: sqlite3.Connection, dossier: Path, code: str) -> Path:
    bloc = db.fetch_one(conn, "SELECT * FROM bloc WHERE code = ?", (code,))
    if bloc is None:
        raise Introuvable(f"Bloc « {code} » introuvable.")
    lignes = db.fetch_all(
        conn, "SELECT * FROM composant WHERE bloc_code = ? AND archive = 0 ORDER BY id", (code,)
    )
    titre = f"Modèle du bloc fonctionnel {code} — {bloc['nom']}"
    return _enregistrer(_construire(conn, lignes, "bloc", code, titre), dossier, "bloc", code)


def generate_modele_ensemble(conn: sqlite3.Connection, dossier: Path, code: str) -> Path:
    ensemble = db.fetch_one(conn, "SELECT * FROM ensemble WHERE code = ? AND archive = 0", (code,))
    if ensemble is None:
        raise Introuvable(f"Ensemble « {code} » introuvable ou archivé.")
    lignes = db.fetch_all(
        conn,
        "SELECT c.*, a.qte AS qte_ensemble FROM affectation a"
        " JOIN composant c ON c.id = a.composant_id"
        " WHERE a.ensemble_code = ? AND c.archive = 0 ORDER BY c.id",
        (code,),
    )
    titre = f"Modèle de l'ensemble {code} — {ensemble['nom']}"
    return _enregistrer(
        _construire(conn, lignes, "ensemble", code, titre), dossier, "ensemble", code
    )
