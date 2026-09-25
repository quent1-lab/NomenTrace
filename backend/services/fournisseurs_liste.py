"""Comparaison des fournisseurs avec une liste Excel (liste des fournisseurs autorisés).

Rien n'est écrit à la comparaison : l'interface renvoie ensuite les créations et les
compléments retenus, appliqués en une seule transaction.
"""

import io
import re
import sqlite3
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend import db
from backend.erreurs import Conflit, ErreurMetier, Introuvable
from backend.services import journal
from backend.services.fournisseurs import CHAMPS_MODIFIABLES
from backend.services.import_colonnes import normaliser_entete
from backend.services.import_lecture import texte_cellule

# En-têtes reconnus (normalisés) ; les autres colonnes, dont la sous-catégorie, sont ignorées.
ENTETES: dict[str, str] = {
    "nom": "nom",
    "fournisseur": "nom",
    "nomfournisseur": "nom",
    "raisonsociale": "nom",
    "categorie": "categorie",
    "famille": "categorie",
    "siteweb": "site_web",
    "site": "site_web",
    "siteinternet": "site_web",
    "url": "site_web",
    "contact": "contact",
    "contactdevis": "contact",
    "email": "contact",
    "mail": "contact",
    "courriel": "contact",
    "ncompte": "numero_compte",
    "nocompte": "numero_compte",
    "numerocompte": "numero_compte",
    "numerodecompte": "numero_compte",
    "compte": "numero_compte",
    "compteclient": "numero_compte",
    "pays": "pays",
    "type": "type",
    "delai": "delai_moyen_j",
    "delaimoyen": "delai_moyen_j",
    "commentaire": "commentaire",
}
CHAMPS_COMPARES: tuple[str, ...] = (
    "categorie",
    "contact",
    "numero_compte",
    "site_web",
    "pays",
    "type",
    "delai_moyen_j",
    "commentaire",
)
LIGNES_ENTETE_MAX = 20


def cle_nom(nom: str) -> str:
    """Nom comparable : casse, accents, espaces et ponctuation ignorés."""
    return normaliser_entete(nom)


def _mots(nom: str) -> set[str]:
    """Mots du nom, normalisés : « RS Components » donne {"rs", "components"}."""
    return {normaliser_entete(m) for m in re.split(r"[\s/(),.\-]+", nom)} - {""}


def _site(texte: str) -> tuple[str, list[str]]:
    """Première adresse, complétée en https:// ; les suivantes sont renvoyées à part."""
    adresses = [ligne.strip() for ligne in texte.splitlines() if ligne.strip()]
    completes = [a if re.match(r"^https?://", a, re.I) else f"https://{a}" for a in adresses]
    return completes[0], completes[1:]


def _valeur(champ: str, brut: Any) -> Any:
    texte = texte_cellule(brut)
    if texte is None:
        return None
    if champ in ("contact", "site_web"):
        # Plusieurs lignes possibles : plusieurs contacts, ou plusieurs sites.
        return "\n".join(ligne.strip() for ligne in texte.splitlines() if ligne.strip())
    if champ == "delai_moyen_j":
        return (
            int(float(texte.replace(",", "."))) if re.fullmatch(r"\d+([.,]\d+)?", texte) else None
        )
    return re.sub(r"\s+", " ", texte)


def _positions(feuille: Any) -> tuple[int, dict[str, int]] | None:
    """Ligne d'en-tête (celle qui contient une colonne de nom) et position des champs."""
    for numero, ligne in enumerate(feuille.iter_rows(values_only=True, max_row=LIGNES_ENTETE_MAX)):
        positions: dict[str, int] = {}
        for index, entete in enumerate(ligne):
            champ = ENTETES.get(normaliser_entete(entete)) if entete else None
            if champ and champ not in positions:
                positions[champ] = index
        if "nom" in positions:
            return numero, positions
    return None


def _fusionner(liste: dict[str, dict], lue: dict[str, Any]) -> None:
    """Un même fournisseur cité deux fois : catégories cumulées, premier renseignement gardé."""
    existant = liste.setdefault(cle_nom(lue["nom"]), {"nom": lue["nom"]})
    for champ, valeur in lue.items():
        if valeur is None or champ == "nom":
            continue
        if champ == "categorie" and existant.get("categorie"):
            categories = existant["categorie"].split(" / ")
            if valeur not in categories:
                existant["categorie"] = " / ".join([*categories, valeur])
        elif not existant.get(champ):
            existant[champ] = valeur


def read_liste(contenu: bytes) -> list[dict]:
    """Lit la liste : la première feuille dont une ligne d'en-tête contient un nom."""
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as erreur:
        raise ErreurMetier(f"Fichier illisible comme classeur Excel ({erreur}).") from erreur
    for feuille in classeur.worksheets:
        trouve = _positions(feuille)
        if trouve is None:
            continue
        ligne_entete, positions = trouve
        liste: dict[str, dict] = {}
        for ligne in feuille.iter_rows(values_only=True, min_row=ligne_entete + 2):
            brut = {c: ligne[i] if i < len(ligne) else None for c, i in positions.items()}
            lue = {champ: _valeur(champ, valeur) for champ, valeur in brut.items()}
            if not lue.get("nom"):
                continue
            if lue.get("site_web"):
                lue["site_web"], autres = _site(lue["site_web"])
                if autres:
                    lue["commentaire"] = "Autres sites : " + ", ".join(autres)
            _fusionner(liste, lue)
        return list(liste.values())
    raise ErreurMetier("Aucune colonne « Nom » ou « Fournisseur » trouvée dans le classeur.")


def _comparable(champ: str, valeur: Any) -> str:
    texte = "" if valeur is None else str(valeur).strip().lower()
    if champ == "site_web":
        texte = re.sub(r"^https?://(www\.)?", "", texte).rstrip("/")
    return re.sub(r"\s+", " ", texte)


def _differences(actuel: dict, lu: dict) -> dict[str, dict]:
    return {
        champ: {"actuel": actuel[champ], "liste": lu[champ]}
        for champ in CHAMPS_COMPARES
        if lu.get(champ) is not None
        and _comparable(champ, actuel[champ]) != _comparable(champ, lu[champ])
    }


def _probable(lu: dict, existants: list[dict]) -> dict | None:
    """Un existant dont le nom est le début de celui de la liste, ou l'inverse (Conrad)."""
    cle, mots = cle_nom(lu["nom"]), _mots(lu["nom"])
    for existant in existants:
        autre = cle_nom(existant["nom"])
        if min(len(cle), len(autre)) < 4:
            continue
        if cle.startswith(autre) or autre.startswith(cle) or _mots(existant["nom"]) <= mots:
            return existant
    return None


def _entree(existant: dict, lu: dict) -> dict:
    return {
        "nom": existant["nom"],
        "liste": lu,
        "archive": bool(existant["archive"]),
        "a_valider": existant["statut"] == "A valider",
        "differences": _differences(existant, lu),
    }


def compare_liste(conn: sqlite3.Connection, lus: list[dict]) -> dict:
    """Classe la liste : à créer, correspondance probable, présents (avec ou sans écarts).

    Les noms identiques (à la casse et aux espaces près) sont appariés d'abord ; seuls les
    fournisseurs restants servent aux correspondances probables.
    """
    existants = db.fetch_all(conn, "SELECT * FROM fournisseur ORDER BY nom COLLATE NOCASE")
    par_cle = {cle_nom(f["nom"]): f for f in existants}
    resultat: dict[str, list] = {"nouveaux": [], "probables": [], "presents": [], "absents": []}
    restants = []
    for lu in sorted(lus, key=lambda f: cle_nom(f["nom"])):
        existant = par_cle.pop(cle_nom(lu["nom"]), None)
        if existant is None:
            restants.append(lu)
        else:
            resultat["presents"].append(_entree(existant, lu))
    libres = list(par_cle.values())
    for lu in restants:
        existant = _probable(lu, libres)
        if existant is None:
            resultat["nouveaux"].append(lu)
        else:
            libres.remove(existant)
            resultat["probables"].append(_entree(existant, lu))
    resultat["absents"] = [f["nom"] for f in libres if not f["archive"]]
    return resultat


def apply_liste(
    conn: sqlite3.Connection, creations: list[dict], completions: list[dict]
) -> dict[str, int]:
    """Crée les fournisseurs retenus et complète les existants, en une seule transaction."""
    with db.transaction(conn):
        for valeurs in creations:
            if db.fetch_one(conn, "SELECT 1 FROM fournisseur WHERE nom = ?", (valeurs["nom"],)):
                raise Conflit(f"Le fournisseur « {valeurs['nom']} » existe déjà.")
            db.insert_row(conn, "fournisseur", valeurs)
            journal.write_journal(conn, "fournisseur", valeurs["nom"], "creation", None, "créé")
        for completion in completions:
            nom = completion["nom"]
            if db.fetch_one(conn, "SELECT 1 FROM fournisseur WHERE nom = ?", (nom,)) is None:
                raise Introuvable(f"Fournisseur « {nom} » introuvable.")
            journal.update_with_journal(
                conn,
                "fournisseur",
                nom,
                {**completion["champs"], "archive": 0},
                CHAMPS_MODIFIABLES | {"archive"},
            )
    return {"crees": len(creations), "completes": len(completions)}
