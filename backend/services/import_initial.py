"""Import initial du fichier de départ, exécuté une seule fois sur une base vide.

C'est le seul module, avec les tests, autorisé à connaître le projet SPOC : il fixe les
paramètres de cette première instance.
"""

import logging
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend import db

journal_log = logging.getLogger(__name__)

# Paramètres de l'instance SPOC. Le budget n'est donné qu'en texte dans _LISEZMOI.
PARAMETRES_INITIAUX: dict[str, str] = {
    "nom_projet": "SPOC",
    "prefixe_id": "SPOC",
    "budget_ht": "3000",
    "taux_tva_defaut": "0.2",
}

# Vocabulaire propre à l'instance SPOC, ajouté aux listes génériques (liste, code, libellé,
# sens pour les types de mouvement).
LISTES_INITIALES: tuple[tuple[str, str, str, str | None], ...] = (
    ("mode_appro", "Stock ecole", "Stock école", None),
    ("mode_appro", "Fourni PFM", "Fourni PFM", None),
    ("mode_appro", "Fourni CEA", "Fourni CEA", None),
    ("mode_appro", "Fabrication PFM", "Fabrication PFM", None),
    ("statut_appro", "Stock-PFM", "Stock PFM", None),
    ("type_mouvement", "Pret ecole", "Prêt école", "Entree"),
    ("type_mouvement", "Retour ecole", "Retour école", "Sortie"),
)

COLONNES_BLOCS: dict[str, str] = {
    "Nom du bloc": "nom",
    "Code": "code",
    "Description": "description",
    "Responsable": "responsable",
}

COLONNES_FOURNISSEURS: dict[str, str] = {
    "Nom": "nom",
    "Type": "type",
    "Base prix par defaut": "base_prix_defaut",
    "Pays": "pays",
    "Site web": "site_web",
    "Compte ecole": "compte_ecole",
    "Delai moyen (j)": "delai_moyen_j",
    "Commentaire": "commentaire",
}

COLONNES_COMPOSANTS: dict[str, str] = {
    "ID": "id",
    "Bloc": "bloc_nom",
    "Fonction": "fonction",
    "Designation": "designation",
    "Ref fabricant": "ref_fabricant",
    "Fabricant": "fabricant",
    "Mode appro": "mode_appro",
    "Fournisseur privilegie": "fournisseur_nom",
    "Lien produit": "lien_produit",
    "Qte besoin": "qte_besoin",
    "Qte rechange": "qte_rechange",
    "Qte dispo ecole": "qte_disponible",
    "PU releve": "pu_releve",
    "Base prix releve": "base_prix_releve",
    "Taux TVA": "taux_tva",
    "Statut choix": "statut_choix",
    "Statut appro": "statut_appro",
    "Criticite": "criticite",
    "Origine exigence": "origine_exigence",
    "Note technique": "note_technique",
}

COLONNES_QUANTITE: tuple[str, ...] = ("qte_besoin", "qte_rechange", "qte_disponible")

# Feuilles de suivi : seules des lignes EXEMPLE y sont attendues, jamais importées.
FEUILLES_EXEMPLE: dict[str, str] = {
    "Commandes": "N commande",
    "Lignes de commande": "ID ligne",
    "Mouvements de stock": "ID mouvement",
}


class ErreurImport(Exception):
    """Fichier de départ inexploitable ; l'import est annulé sans rien écrire."""


def _nettoyer(valeur: Any) -> Any:
    """Supprime les espaces de bord ; une chaîne vide devient None."""
    if isinstance(valeur, str):
        valeur = valeur.strip()
        return valeur or None
    return valeur


def read_sheet(classeur: Any, nom_feuille: str, colonnes: dict[str, str]) -> list[dict]:
    """Lit une feuille en dictionnaires selon la correspondance en-tête -> colonne."""
    if nom_feuille not in classeur.sheetnames:
        raise ErreurImport(f"Feuille « {nom_feuille} » absente du fichier.")
    feuille: Worksheet = classeur[nom_feuille]
    lignes = feuille.iter_rows(values_only=True)
    entetes = [_nettoyer(v) for v in next(lignes, ())]
    manquants = [e for e in colonnes if e not in entetes]
    if manquants:
        raise ErreurImport(
            f"Feuille « {nom_feuille} » : colonnes manquantes : {', '.join(manquants)}."
        )
    positions = {colonnes[e]: entetes.index(e) for e in colonnes}
    resultat = []
    for ligne in lignes:
        enregistrement = {
            champ: _nettoyer(ligne[i]) if i < len(ligne) else None for champ, i in positions.items()
        }
        if any(v is not None for v in enregistrement.values()):
            resultat.append(enregistrement)
    return resultat


def _entier(valeur: Any, contexte: str) -> int:
    """Convertit une quantité en entier ; vide vaut 0, une fraction est refusée."""
    if valeur is None:
        return 0
    if isinstance(valeur, bool) or not isinstance(valeur, int | float) or valeur != int(valeur):
        raise ErreurImport(f"{contexte} : quantité non entière « {valeur} ».")
    return int(valeur)


def _prepare_composants(composants: list[dict], codes_par_nom: dict[str, str]) -> list[dict]:
    """Remplace le nom du bloc par son code et convertit les quantités."""
    prepares = []
    for c in composants:
        nom_bloc = c.pop("bloc_nom")
        if nom_bloc not in codes_par_nom:
            raise ErreurImport(f"Composant {c['id']} : bloc « {nom_bloc} » inconnu.")
        c["bloc_code"] = codes_par_nom[nom_bloc]
        for champ in COLONNES_QUANTITE:
            c[champ] = _entier(c[champ], f"Composant {c['id']}, {champ}")
        if c["taux_tva"] is None:
            c["taux_tva"] = float(PARAMETRES_INITIAUX["taux_tva_defaut"])
        prepares.append(c)
    return prepares


def _check_feuilles_exemple(classeur: Any) -> None:
    """Vérifie que les feuilles de suivi ne contiennent que des lignes EXEMPLE."""
    for nom_feuille, colonne_id in FEUILLES_EXEMPLE.items():
        for ligne in read_sheet(classeur, nom_feuille, {colonne_id: "id"}):
            identifiant = str(ligne["id"] or "")
            if not identifiant.startswith("EXEMPLE"):
                raise ErreurImport(
                    f"Feuille « {nom_feuille} » : la ligne « {identifiant} » n'est pas un "
                    "EXEMPLE ; l'import initial ne prend pas en charge le suivi réel."
                )


def read_fichier(chemin: Path) -> dict[str, list[dict]]:
    """Lit et valide tout le fichier de départ, sans rien écrire en base."""
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    try:
        blocs = read_sheet(classeur, "Blocs", COLONNES_BLOCS)
        fournisseurs = read_sheet(classeur, "Fournisseurs", COLONNES_FOURNISSEURS)
        composants = read_sheet(classeur, "Composants", COLONNES_COMPOSANTS)
        _check_feuilles_exemple(classeur)
    finally:
        classeur.close()
    for ordre, bloc in enumerate(blocs, start=1):
        bloc["ordre"] = ordre
    for fournisseur in fournisseurs:
        fournisseur["delai_moyen_j"] = (
            _entier(fournisseur["delai_moyen_j"], f"Fournisseur {fournisseur['nom']}")
            if fournisseur["delai_moyen_j"] is not None
            else None
        )
    codes_par_nom = {b["nom"]: b["code"] for b in blocs}
    return {
        "bloc": blocs,
        "fournisseur": fournisseurs,
        "composant": _prepare_composants(composants, codes_par_nom),
    }


def _insert_rows(conn: sqlite3.Connection, table: str, lignes: list[dict]) -> None:
    """Insère des lignes ; les noms de table et de colonnes viennent des constantes ci-dessus."""
    if not lignes:
        return
    colonnes = list(lignes[0])
    marques = ", ".join("?" for _ in colonnes)
    sql = f"INSERT INTO {table} ({', '.join(colonnes)}) VALUES ({marques})"  # noqa: S608
    conn.executemany(sql, [tuple(ligne[c] for c in colonnes) for ligne in lignes])


def _insert_listes(conn: sqlite3.Connection) -> None:
    """Ajoute le vocabulaire de l'instance à la fin de chaque liste, sans doublon."""
    for liste, code, libelle, sens in LISTES_INITIALES:
        conn.execute(
            "INSERT OR IGNORE INTO valeur_liste (liste, code, libelle, ordre, sens)"
            " SELECT ?, ?, ?, COALESCE(MAX(ordre), 0) + 1, ? FROM valeur_liste WHERE liste = ?",
            (liste, code, libelle, sens, liste),
        )


def run_import_initial(conn: sqlite3.Connection, chemin: Path) -> bool:
    """Importe le fichier de départ si la base est vide. Renvoie True si l'import a eu lieu."""
    deja = conn.execute("SELECT COUNT(*) FROM composant").fetchone()[0]
    if deja:
        journal_log.info("Import initial ignoré : la base contient déjà %d composants.", deja)
        return False
    if not chemin.exists():
        journal_log.warning("Import initial ignoré : fichier %s introuvable.", chemin)
        return False
    donnees = read_fichier(chemin)
    maintenant = datetime.now().isoformat(timespec="seconds")
    for composant in donnees["composant"]:
        composant["cree_le"] = maintenant
        composant["modifie_le"] = maintenant
    parametres = [{"cle": c, "valeur": v} for c, v in PARAMETRES_INITIAUX.items()]
    synthese = ", ".join(f"{table} : {len(lignes)}" for table, lignes in donnees.items())
    with db.transaction(conn, immediate=True):
        _insert_rows(conn, "parametre", parametres)
        _insert_listes(conn)
        for table in ("bloc", "fournisseur", "composant"):
            _insert_rows(conn, table, donnees[table])
        conn.execute(
            "INSERT INTO journal (horodatage, table_cible, cle_cible, champ, nouvelle_valeur,"
            " origine) VALUES (?, '*', ?, 'import_initial', ?, 'import')",
            (maintenant, chemin.name, synthese),
        )
    journal_log.info("Import initial terminé depuis %s — %s.", chemin.name, synthese)
    return True
