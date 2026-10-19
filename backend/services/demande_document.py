"""Document d'une demande de devis à envoyer au fournisseur, en Excel et en PDF.

Il reprend l'établissement (Paramètres › Projet) et son logo, le demandeur, le fournisseur
et, pour chaque ligne de la commande, la référence, la désignation, le fabricant, la
quantité et le lien produit. Aucun prix : c'est au fournisseur de chiffrer.
"""

import io
import re
import sqlite3
import unicodedata
from datetime import date

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImageExcel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend import db
from backend.services import commandes, logo, parametres, tableur
from backend.services.droits import Utilisateur

PHRASE = "Merci de nous adresser un devis HT pour les articles ci-dessous."
COLONNES: tuple[tuple[str, str], ...] = (
    ("ref_fabricant", "Réf. fabricant"),
    ("designation", "Désignation"),
    ("fabricant", "Fabricant"),
    ("qte", "Qté"),
    ("lien_produit", "Lien produit"),
)


def _jjmmaaaa(iso: str | None) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}" if iso else ""


def demandeur(conn: sqlite3.Connection, utilisateur: Utilisateur | None) -> dict:
    """La personne connectée ; en mode local, faute de compte, le contact achats."""
    if utilisateur is not None and not utilisateur.local:
        return {"nom": utilisateur.nom, "mail": utilisateur.identifiant}
    return {
        "nom": parametres.get_parametre(conn, "contact_achats_nom"),
        "mail": parametres.get_parametre(conn, "contact_achats_mail"),
    }


def get_donnees(
    conn: sqlite3.Connection, numero: str, demandeur: dict, reponse_avant: str | None
) -> dict:
    """Tout ce que le document affiche, lu en base."""
    commande = commandes.get_commande(conn, numero)
    fournisseur = (
        db.fetch_one(
            conn,
            "SELECT nom, numero_compte FROM fournisseur WHERE nom = ?",
            (commande["fournisseur_nom"],),
        )
        or {}
    )
    lignes = db.fetch_all(
        conn,
        "SELECT c.ref_fabricant, c.designation, c.fabricant, l.qte_commandee AS qte,"
        " c.lien_produit FROM ligne_commande l JOIN composant c ON c.id = l.composant_id"
        " WHERE l.commande_numero = ? AND l.statut_ligne <> 'Annulee' ORDER BY l.id",
        (numero,),
    )
    reglages = parametres.list_parametres(conn)
    return {
        "numero": numero,
        "date": date.today().isoformat(),
        "reponse_avant": reponse_avant,
        "projet": reglages.get("nom_projet") or "",
        "etablissement": {
            "nom": reglages.get("etablissement_nom") or "",
            "livraison": reglages.get("adresse_livraison") or "",
            "facturation": reglages.get("adresse_facturation") or "",
            "contact": reglages.get("contact_achats_nom") or "",
            "mail": reglages.get("contact_achats_mail") or "",
            "tel": reglages.get("contact_achats_tel") or "",
        },
        "demandeur": demandeur,
        "fournisseur": {
            "nom": commande["fournisseur_nom"] or "",
            "numero_compte": fournisseur.get("numero_compte") or "",
        },
        "reference_externe": commande["reference_externe"] or "",
        "lignes": lignes,
        "logo": logo.get_logo(conn),
    }


def _taille_logo(image: dict, largeur_max: float, hauteur_max: float) -> tuple[float, float]:
    """Dimensions du logo dans la boîte permise, proportions gardées."""
    echelle = min(largeur_max / image["largeur"], hauteur_max / image["hauteur"])
    return image["largeur"] * echelle, image["hauteur"] * echelle


def nom_fichier(donnees: dict, extension: str) -> str:
    """« Demande de devis CMD-012 Mouser.pdf », sans caractère gênant pour un fichier."""
    base = f"Demande de devis {donnees['numero']} {donnees['fournisseur']['nom']}".strip()
    return re.sub(r"[^\w .-]+", "_", base) + extension


def _entete(donnees: dict) -> list[tuple[str, str]]:
    """Lignes « libellé : valeur » de l'en-tête, vides écartées."""
    e, d, f = donnees["etablissement"], donnees["demandeur"], donnees["fournisseur"]
    contact = " · ".join(v for v in (e["contact"], e["mail"], e["tel"]) if v)
    paires = [
        ("Établissement", e["nom"]),
        ("Projet", donnees["projet"]),
        ("Demande n°", donnees["numero"]),
        ("Date", _jjmmaaaa(donnees["date"])),
        ("Demandeur", " · ".join(v for v in (d.get("nom"), d.get("mail")) if v)),
        ("Contact achats", contact),
        ("Fournisseur", f["nom"]),
        ("N° de compte client", f["numero_compte"]),
        ("Votre référence", donnees["reference_externe"]),
        ("Réponse souhaitée avant le", _jjmmaaaa(donnees["reponse_avant"])),
        ("Adresse de livraison", e["livraison"]),
        ("Adresse de facturation", e["facturation"]),
    ]
    return [(libelle, valeur) for libelle, valeur in paires if valeur]


# --- Excel --------------------------------------------------------------------------------------


def build_excel(donnees: dict) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Demande de devis"
    feuille.append(["Demande de devis"])
    feuille["A1"].font = Font(bold=True, size=14)
    feuille.append([])
    for libelle, valeur in _entete(donnees):
        feuille.append([libelle, valeur])
        feuille.cell(feuille.max_row, 1).font = Font(bold=True)
        feuille.cell(feuille.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    feuille.append([])
    feuille.append([PHRASE])
    feuille.append([])
    feuille.append([titre for _, titre in COLONNES])
    fond = PatternFill("solid", fgColor="E6ECF3")
    trait = Side(style="thin", color="D9DEE5")
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
        cellule.fill = fond
    for ligne in donnees["lignes"]:
        feuille.append([ligne[cle] for cle, _ in COLONNES])
        for cellule in feuille[feuille.max_row]:
            cellule.border = Border(bottom=trait)
            cellule.alignment = Alignment(wrap_text=True, vertical="top")
        lien = ligne["lien_produit"]
        if lien and lien.startswith(("http://", "https://")):
            feuille.cell(feuille.max_row, len(COLONNES)).hyperlink = lien
    for colonne, largeur in zip("ABCDE", (22, 50, 22, 8, 45), strict=True):
        feuille.column_dimensions[colonne].width = largeur
    if donnees.get("logo"):
        # En haut à droite, dans la colonne du lien, à côté de l'en-tête.
        image = ImageExcel(io.BytesIO(donnees["logo"]["contenu"]))
        image.width, image.height = _taille_logo(donnees["logo"], 300, 70)
        feuille.add_image(image, "E1")
    tableur.neutraliser_formules(classeur)
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


# --- PDF ----------------------------------------------------------------------------------------

# Les polices intégrées du PDF ne couvrent que le latin-1 : les quelques caractères typographiques
# qui en sortent sont remplacés par leur équivalent le plus proche.
REMPLACEMENTS = str.maketrans(
    {
        "œ": "oe",
        "Œ": "OE",
        "’": "'",
        "‘": "'",
        "“": '"',
        "”": '"',
        "–": "-",
        "—": "-",
        "…": "...",
        " ": " ",
        " ": " ",
        "€": "EUR",
    }
)


def _latin1(texte: object) -> str:
    propre = unicodedata.normalize("NFC", str(texte or "")).translate(REMPLACEMENTS)
    return propre.encode("latin-1", errors="replace").decode("latin-1")


def _lien_court(lien: str | None) -> str:
    """Domaine seul, pour tenir dans la colonne ; le lien complet reste cliquable."""
    if not lien:
        return ""
    return re.sub(r"^https?://(www\.)?", "", lien).split("/")[0]


def build_pdf(donnees: dict) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_title(_latin1(f"Demande de devis {donnees['numero']}"))
    pdf.add_page()
    bas_logo = 0.0
    if donnees.get("logo"):
        # En haut à droite, au plus 50 × 22 mm ; l'en-tête commence sous le logo.
        largeur, hauteur = _taille_logo(donnees["logo"], 50, 22)
        pdf.image(
            io.BytesIO(donnees["logo"]["contenu"]),
            x=pdf.w - pdf.r_margin - largeur,
            y=pdf.t_margin,
            w=largeur,
            h=hauteur,
        )
        bas_logo = pdf.t_margin + hauteur
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 10, _latin1("Demande de devis"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_y(max(pdf.get_y(), bas_logo + 2))
    pdf.set_font("helvetica", size=9)
    for libelle, valeur in _entete(donnees):
        pdf.set_font("helvetica", "B", 9)
        pdf.cell(48, 5, _latin1(libelle))
        pdf.set_font("helvetica", size=9)
        pdf.multi_cell(0, 5, _latin1(valeur), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("helvetica", size=10)
    pdf.multi_cell(0, 6, _latin1(PHRASE), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("helvetica", size=8.5)
    entete = FontFace(emphasis="BOLD", fill_color=(230, 236, 243))
    with pdf.table(
        col_widths=(30, 70, 30, 12, 40),
        text_align=("LEFT", "LEFT", "LEFT", "RIGHT", "LEFT"),
        headings_style=entete,
        line_height=4.5,
    ) as table:
        table.row([_latin1(titre) for _, titre in COLONNES])
        for ligne in donnees["lignes"]:
            rang = table.row()
            for cle, _ in COLONNES[:-1]:
                rang.cell(_latin1(ligne[cle] if ligne[cle] is not None else ""))
            lien = ligne["lien_produit"] or ""
            cliquable = lien.startswith(("http://", "https://"))
            rang.cell(_latin1(_lien_court(lien)), link=lien if cliquable else None)
    return bytes(pdf.output())
