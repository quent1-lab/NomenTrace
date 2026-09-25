"""Export Excel automatique de la base : écriture atomique, anti-rafale, fichier verrouillé."""

import io
import logging
import os
import sqlite3
import tempfile
import threading
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from backend import db
from backend.arrondi import round_value
from backend.services import attributs, tableur

journal_log = logging.getLogger(__name__)

NOM_FICHIER: str = "nomenclature.xlsx"
AVERTISSEMENT: str = (
    "Fichier généré automatiquement par Nomentrace à partir de la base. "
    "Toute modification manuelle sera écrasée au prochain export."
)
# Vues de lecture d'abord, puis une feuille par table. Noms fixés dans le code.
FEUILLES_VUES: tuple[tuple[str, str], ...] = (
    ("Blocs", "SELECT * FROM v_bloc ORDER BY ordre, code"),
    ("Ensembles", "SELECT * FROM v_ensemble ORDER BY ordre, code"),
    ("Affectations", "SELECT * FROM v_ensemble_composant ORDER BY ensemble_code, composant_id"),
)
TABLES: tuple[str, ...] = (
    "parametre",
    "valeur_liste",
    "bloc",
    "ensemble",
    "fournisseur",
    "composant",
    "attribut",
    "attribut_valeur",
    "composant_attribut",
    "affectation",
    "commande",
    "ligne_commande",
    "mouvement_stock",
    "document",
    "journal",
)
FORMAT_MONTANT: str = "#,##0.00"


def _write_rows(feuille: Worksheet, lignes: list[dict], colonnes: list[str]) -> None:
    feuille.append(colonnes)
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
    for ligne in lignes:
        feuille.append([round_value(c, ligne[c]) for c in colonnes])
        for cellule, colonne in zip(feuille[feuille.max_row], colonnes, strict=True):
            if isinstance(cellule.value, float) and not colonne.startswith("taux"):
                cellule.number_format = FORMAT_MONTANT


def _write_attributs(feuille: Worksheet, conn: sqlite3.Connection) -> None:
    """Une colonne par attribut actif à droite de la feuille composant, valeurs lisibles."""
    colonne_id = next(c.column for c in feuille[1] if c.value == "id")
    valeurs = attributs.valeurs_par_composant(conn)
    libelles = attributs.libelles_valeurs(conn)
    for attribut in attributs.list_attributs(conn, actifs_seulement=True):
        colonne = feuille.max_column + 1
        feuille.cell(1, colonne, attributs.entete(attribut)).font = Font(bold=True)
        for ligne in range(2, feuille.max_row + 1):
            identifiant = feuille.cell(ligne, colonne_id).value
            valeur = valeurs.get(identifiant, {}).get(attribut["code"])
            feuille.cell(ligne, colonne, attributs.affichage(attribut, valeur, libelles))


def _colonnes(conn: sqlite3.Connection, sql: str) -> list[str]:
    curseur = conn.execute(sql + " LIMIT 0")
    return [d[0] for d in curseur.description]


def build_workbook(chemin_base: Path) -> Workbook:
    """Construit le classeur complet à partir d'une connexion de lecture dédiée."""
    conn = db.connect(chemin_base)
    try:
        classeur = Workbook()
        pilotage = classeur.active
        pilotage.title = "Pilotage"
        pilotage.append([AVERTISSEMENT])
        pilotage["A1"].font = Font(bold=True, color="C53030")
        pilotage.append([])
        pilotage.append(["Indicateur", "Valeur"])
        for cle, valeur in (db.fetch_one(conn, "SELECT * FROM v_pilotage") or {}).items():
            pilotage.append([cle, round_value(cle, valeur)])
        requetes = list(FEUILLES_VUES) + [
            (table, f"SELECT * FROM {table}")  # noqa: S608
            for table in TABLES
        ]
        for nom, sql in requetes:
            feuille = classeur.create_sheet(nom)
            _write_rows(feuille, db.fetch_all(conn, sql), _colonnes(conn, sql))
            if nom == "composant":
                _write_attributs(feuille, conn)
        tableur.neutraliser_formules(classeur)
        return classeur
    finally:
        conn.close()


def build_export_bytes(chemin_base: Path) -> bytes:
    """Classeur complet régénéré à la demande, en mémoire, pour un téléchargement."""
    tampon = io.BytesIO()
    build_workbook(chemin_base).save(tampon)
    return tampon.getvalue()


def nom_telechargement(maintenant: datetime | None = None) -> str:
    """nomentrace_export_AAAAMMJJ_HHMM.xlsx"""
    return f"nomentrace_export_{(maintenant or datetime.now()):%Y%m%d_%H%M}.xlsx"


def write_export(chemin_base: Path, dossier: Path) -> Path:
    """Écrit l'export de façon atomique : fichier temporaire puis os.replace().

    Lève PermissionError si le fichier cible est ouvert dans Excel.
    """
    dossier.mkdir(parents=True, exist_ok=True)
    cible = dossier / NOM_FICHIER
    descripteur, temporaire = tempfile.mkstemp(prefix="export_", suffix=".tmp", dir=dossier)
    os.close(descripteur)
    try:
        build_workbook(chemin_base).save(temporaire)
        os.replace(temporaire, cible)
    finally:
        if os.path.exists(temporaire):
            os.remove(temporaire)
    return cible


class PlanificateurExport:
    """Déclenche l'export au plus tôt `delai` secondes après la dernière modification.

    Un seul fil de travail. Si le fichier est verrouillé, l'export reste en attente et
    est retenté toutes les `reessai` secondes.
    """

    def __init__(
        self, chemin_base: Path, dossier: Path, delai: float = 2.0, reessai: float = 30.0
    ) -> None:
        self.chemin_base = chemin_base
        self.dossier = dossier
        self.delai = delai
        self.reessai = reessai
        self.en_attente = False
        self._echeance: float | None = None
        self._verrou = threading.Lock()
        self._ecriture = threading.Lock()
        self._reveil = threading.Event()
        self._arret = threading.Event()
        self._fil = threading.Thread(target=self._boucle, name="export-excel", daemon=True)

    def start(self) -> None:
        self._fil.start()

    def stop(self) -> None:
        self._arret.set()
        self._reveil.set()
        self._fil.join(timeout=5)

    def signaler(self) -> None:
        """Signale une modification : l'export est repoussé de `delai` secondes."""
        with self._verrou:
            self._echeance = time.monotonic() + self.delai
        self._reveil.set()

    def export_now(self) -> bool:
        """Exporte immédiatement. Renvoie False si le fichier est verrouillé."""
        with self._verrou:
            self._echeance = None
        return self._tenter()

    def _tenter(self) -> bool:
        try:
            with self._ecriture:
                write_export(self.chemin_base, self.dossier)
        except PermissionError as erreur:
            journal_log.warning("Export en attente, fichier verrouillé (Excel ?) : %s", erreur)
            with self._verrou:
                self.en_attente = True
                self._echeance = time.monotonic() + self.reessai
            self._reveil.set()
            return False
        except OSError as erreur:
            journal_log.error("Export impossible : %s", erreur)
            with self._verrou:
                self.en_attente = True
                self._echeance = time.monotonic() + self.reessai
            self._reveil.set()
            return False
        with self._verrou:
            self.en_attente = False
        journal_log.info("Export Excel à jour.")
        return True

    def _boucle(self) -> None:
        while not self._arret.is_set():
            with self._verrou:
                echeance = self._echeance
            attente = None if echeance is None else max(0.0, echeance - time.monotonic())
            self._reveil.wait(timeout=attente)
            self._reveil.clear()
            if self._arret.is_set():
                return
            with self._verrou:
                pret = self._echeance is not None and time.monotonic() >= self._echeance
                if pret:
                    self._echeance = None
            if pret:
                self._tenter()
