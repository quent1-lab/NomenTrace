"""Export de la liste des composants telle qu'affichée : mêmes filtres, tri et colonnes."""

import io
import sqlite3
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend import db
from backend.arrondi import round_value
from backend.erreurs import ErreurMetier
from backend.services import composants
from backend.services.composants import FiltresComposants

FORMAT_ENTIER: str = "#,##0"
FORMAT_MONTANT: str = '#,##0.00 "€"'


@dataclass(frozen=True)
class ColonneExport:
    """Colonne exportable : en-tête affiché à l'écran et nature de la valeur."""

    titre: str
    nature: str  # texte | liste | entier | montant | pu_releve | avancement


# Liste blanche des colonnes exportables, avec les en-têtes de l'écran composants.
COLONNES: dict[str, ColonneExport] = {
    "id": ColonneExport("ID", "texte"),
    "bloc_code": ColonneExport("Bloc", "texte"),
    "fonction": ColonneExport("Fonction", "texte"),
    "designation": ColonneExport("Désignation", "texte"),
    "lien_produit": ColonneExport("Page produit", "texte"),
    "ref_fabricant": ColonneExport("Réf fabricant", "texte"),
    "mode_appro": ColonneExport("Mode appro", "liste"),
    "fournisseur_nom": ColonneExport("Fournisseur", "texte"),
    "qte_besoin": ColonneExport("Besoin", "entier"),
    "qte_rechange": ColonneExport("Rech.", "entier"),
    "qte_disponible": ColonneExport("Dispo.", "entier"),
    "qte_a_acheter": ColonneExport("À acheter", "entier"),
    "qte_affectee": ColonneExport("Qté affectée", "entier"),
    "pu_releve": ColonneExport("PU relevé", "pu_releve"),
    "pu_ht": ColonneExport("PU HT", "montant"),
    "total_ht": ColonneExport("Total HT", "montant"),
    "statut_choix": ColonneExport("Statut choix", "liste"),
    "statut_appro": ColonneExport("Statut appro", "liste"),
    "criticite": ColonneExport("Criticité", "liste"),
    "avancement": ColonneExport("Avancement", "avancement"),
}

# L'avancement est calculé par v_composant : liste figée, libellés de format.js.
LIBELLES_AVANCEMENT: dict[str, str] = {
    "A commander": "À commander",
    "Commande": "Commandé",
    "Recu": "Reçu",
    "Hors achat": "Hors achat",
}


def parse_colonnes(texte: str | None) -> list[str]:
    """Colonnes demandées, séparées par des virgules, validées contre la liste blanche."""
    demandees = [c.strip() for c in (texte or "").split(",") if c.strip()]
    if not demandees:
        return list(COLONNES)
    inconnues = [c for c in demandees if c not in COLONNES]
    if inconnues:
        raise ErreurMetier(f"Colonne(s) inconnue(s) : {', '.join(inconnues)}.")
    if len(set(demandees)) != len(demandees):
        raise ErreurMetier("Une colonne est demandée deux fois.")
    return demandees


def _libelles_listes(conn: sqlite3.Connection) -> dict[tuple[str, str], str]:
    lignes = db.fetch_all(conn, "SELECT liste, code, libelle FROM valeur_liste")
    return {(ligne["liste"], ligne["code"]): ligne["libelle"] for ligne in lignes}


def _cellule(cle: str, composant: dict, libelles: dict[tuple[str, str], str]) -> object:
    """Valeur de la cellule : libellé pour les listes, nombre arrondi pour les montants."""
    valeur = composant[cle]
    nature = COLONNES[cle].nature
    if nature == "liste" and valeur is not None:
        return libelles.get((cle, valeur), valeur)
    if nature == "avancement":
        return LIBELLES_AVANCEMENT.get(valeur, valeur)
    if nature == "pu_releve" and valeur is None:
        return "à chiffrer" if composant["mode_appro"] == "Achat" else None
    return round_value(cle, valeur)


def _format(cle: str, composant: dict) -> str | None:
    """Format numérique de la cellule ; le PU relevé porte sa base (HT ou TTC)."""
    nature = COLONNES[cle].nature
    if nature == "entier":
        return FORMAT_ENTIER
    if nature == "montant":
        return FORMAT_MONTANT
    if nature == "pu_releve" and composant["pu_releve"] is not None:
        return f'#,##0.00 "€ {composant["base_prix_releve"]}"'
    return None


def _ligne_totaux(feuille: Worksheet, colonnes: list[str], lignes: list[dict]) -> None:
    """Ligne de totaux comme à l'écran : nombre de composants, total HT sous sa colonne."""
    total = round_value("total_ht", sum(c["total_ht"] or 0.0 for c in lignes))
    valeurs: list[object] = [None] * len(colonnes)
    valeurs[0] = f"{len(lignes)} composant(s) affiché(s)"
    if "total_ht" in colonnes:
        rang = colonnes.index("total_ht")
        valeurs[rang] = total
        if rang > 1:
            valeurs[rang - 1] = "Total HT"
    feuille.append(valeurs)
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
    if "total_ht" in colonnes:
        rang = colonnes.index("total_ht") + 1
        feuille.cell(feuille.max_row, rang).number_format = FORMAT_MONTANT


def build_export_composants(
    conn: sqlite3.Connection, filtres: FiltresComposants, colonnes: list[str]
) -> bytes:
    """Classeur de la liste filtrée et triée, colonnes dans l'ordre demandé, en mémoire."""
    lignes = composants.list_composants(conn, filtres)
    libelles = _libelles_listes(conn)
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Composants"
    feuille.append([COLONNES[c].titre for c in colonnes])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for composant in lignes:
        feuille.append([_cellule(c, composant, libelles) for c in colonnes])
        for rang, cle in enumerate(colonnes, start=1):
            numero = _format(cle, composant)
            if numero:
                feuille.cell(feuille.max_row, rang).number_format = numero
    _ligne_totaux(feuille, colonnes, lignes)
    feuille.freeze_panes = "B2"
    for rang, cle in enumerate(colonnes, start=1):
        largeur = 40 if cle in ("designation", "fonction", "lien_produit") else 14
        feuille.column_dimensions[get_column_letter(rang)].width = largeur
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def nom_fichier(maintenant: datetime | None = None) -> str:
    """nomentrace_composants_AAAAMMJJ_HHMM.xlsx"""
    return f"nomentrace_composants_{(maintenant or datetime.now()):%Y%m%d_%H%M}.xlsx"
