"""Lecture d'un fichier de l'équipe et conversion de ses cellules en valeurs de la base.

Aucune écriture ici : le résultat sert à classer chaque ligne (import_analyse).
"""

import io
import re
import sqlite3
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from backend import db
from backend.services import attributs, listes, parametres
from backend.services.import_colonnes import (
    ALIAS_ENTETES,
    COLONNE_ENSEMBLE,
    COLONNE_QTE_ENSEMBLE,
    COLONNES,
    NOM_FEUILLE,
    NOM_FEUILLE_META,
    OBLIGATOIRES,
    normaliser_entete,
)

TOUTES: tuple = (*COLONNES, COLONNE_QTE_ENSEMBLE, COLONNE_ENSEMBLE)
NATURES: dict[str, str] = {c.champ: c.nature for c in TOUTES}
LIBELLES: dict[str, str] = {c.champ: c.entete for c in TOUTES}


class FichierIllisible(Exception):
    """Le fichier n'est pas un classeur exploitable."""


@dataclass
class Classeur:
    """Contenu utile d'un fichier déposé."""

    meta: dict[str, str]
    colonnes: set[str]
    lignes: list[tuple[int, dict[str, Any]]]  # (numéro de ligne Excel, champ -> valeur brute)
    entetes: dict[str, str]  # champ -> en-tête tel qu'écrit dans le fichier


@dataclass
class Contexte:
    """Référentiels de la base utilisés pour convertir les valeurs."""

    blocs: dict[str, str]
    fournisseurs: dict[str, str]
    listes: dict[str, dict[str, str]]
    ensembles: dict[str, str]
    taux_defaut: float
    bloc_defaut: str | None = None
    attributs: dict[str, dict] = field(default_factory=dict)


@dataclass
class LigneLue:
    """Une ligne convertie : valeurs typées, erreurs bloquantes, avertissements."""

    numero: int
    brut: dict[str, str]
    valeurs: dict[str, Any] = field(default_factory=dict)
    erreurs: list[str] = field(default_factory=list)
    avertissements: list[str] = field(default_factory=list)
    # Fournisseurs, valeurs de liste et ensembles cités mais absents de la base.
    entites: list[dict] = field(default_factory=list)


def texte_cellule(valeur: Any) -> str | None:
    """Texte d'une cellule : les nombres entiers perdent leur « .0 », les dates leur heure."""
    if valeur is None:
        return None
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    if isinstance(valeur, datetime):
        valeur = valeur.date().isoformat()
    texte = str(valeur).strip()
    return texte or None


def alias_attributs(conn: sqlite3.Connection) -> dict[str, str]:
    """En-têtes reconnus pour les attributs actifs : « Tension (V) », « Tension », « tension »."""
    alias = {}
    for attribut in attributs.list_attributs(conn, actifs_seulement=True):
        for texte in (attributs.entete(attribut), attribut["libelle"], attribut["code"]):
            alias.setdefault(normaliser_entete(texte), attributs.champ(attribut["code"]))
    return alias


def read_classeur(contenu: bytes, alias_attributs: dict[str, str] | None = None) -> Classeur:
    """Lit la feuille des composants et, si présente, la feuille d'identification.

    `alias_attributs` ajoute les en-têtes des attributs paramétrables aux en-têtes connus.
    """
    alias = {**(alias_attributs or {}), **ALIAS_ENTETES}
    try:
        classeur = load_workbook(io.BytesIO(contenu), data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as erreur:
        raise FichierIllisible(f"Fichier illisible comme classeur Excel ({erreur}).") from erreur
    meta = {}
    if NOM_FEUILLE_META in classeur.sheetnames:
        for cle, valeur, *_ in classeur[NOM_FEUILLE_META].iter_rows(values_only=True):
            if cle:
                meta[str(cle)] = texte_cellule(valeur) or ""
    feuille = (
        classeur[NOM_FEUILLE] if NOM_FEUILLE in classeur.sheetnames else classeur.worksheets[0]
    )
    lignes = feuille.iter_rows(values_only=True)
    entetes_brutes = next(lignes, ())
    positions, entetes = {}, {}
    for index, entete in enumerate(entetes_brutes):
        champ = alias.get(normaliser_entete(entete)) if entete else None
        if champ and champ not in positions:
            positions[champ] = index
            entetes[champ] = str(entete)
    if "designation" not in positions and "id" not in positions:
        raise FichierIllisible("Aucune colonne « ID » ni « Désignation » : ce n'est pas un modèle.")
    resultat = []
    for numero, ligne in enumerate(lignes, start=2):
        brut = {champ: ligne[i] if i < len(ligne) else None for champ, i in positions.items()}
        resultat.append((numero, brut))
    return Classeur(meta, set(positions), resultat, entetes)


def _cle(texte: str) -> str:
    decompose = unicodedata.normalize("NFKD", texte.lower())
    return re.sub(
        r"\s+", " ", "".join(c for c in decompose if not unicodedata.combining(c))
    ).strip()


def load_contexte(conn: sqlite3.Connection, bloc_defaut: str | None) -> Contexte:
    """Charge blocs, fournisseurs et listes, indexés par code et par libellé normalisés."""
    blocs = {}
    for bloc in db.fetch_all(conn, "SELECT code, nom FROM bloc"):
        blocs[_cle(bloc["code"])] = bloc["code"]
        blocs[_cle(bloc["nom"])] = bloc["code"]
    fournisseurs = {
        _cle(f["nom"]): f["nom"]
        for f in db.fetch_all(conn, "SELECT nom FROM fournisseur WHERE archive = 0")
    }
    listes: dict[str, dict[str, str]] = {}
    for valeur in db.fetch_all(conn, "SELECT liste, code, libelle FROM valeur_liste"):
        index = listes.setdefault(valeur["liste"], {})
        index[_cle(valeur["code"])] = valeur["code"]
        index[_cle(valeur["libelle"])] = valeur["code"]
    ensembles = {}
    for ensemble in db.fetch_all(conn, "SELECT code, nom FROM ensemble WHERE archive = 0"):
        ensembles[_cle(ensemble["code"])] = ensemble["code"]
        ensembles[_cle(ensemble["nom"])] = ensemble["code"]
    taux = parametres.get_taux_tva_defaut(conn)
    definitions = {a["code"]: a for a in attributs.list_attributs(conn, actifs_seulement=True)}
    return Contexte(blocs, fournisseurs, listes, ensembles, taux, bloc_defaut, definitions)


def _nombre(texte: str) -> float | None:
    nettoye = re.sub(r"[\s  €%]", "", texte).replace(",", ".")
    try:
        return float(nettoye)
    except ValueError:
        return None


def code_ensemble(texte: str) -> str:
    """Code proposé pour un nouvel ensemble : majuscules, chiffres et tirets."""
    decompose = unicodedata.normalize("NFKD", texte.upper())
    sans_accents = "".join(c for c in decompose if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]+", "-", sans_accents).strip("-")[:20] or "ENSEMBLE"


def _reference(champ: str, texte: str, ctx: Contexte) -> tuple[Any, str | None, dict | None]:
    """Bloc, fournisseur, valeur de liste ou ensemble : existant, ou proposé à la création."""
    nature, libelle = NATURES[champ], LIBELLES[champ]
    if nature == "bloc":
        code = ctx.blocs.get(_cle(texte))
        return (code, None, None) if code else (None, f"{libelle} : « {texte} » inconnu", None)
    if nature == "fournisseur":
        nom = ctx.fournisseurs.get(_cle(texte))
        return (nom, None, None) if nom else (texte, None, {"type": "fournisseur", "nom": texte})
    if nature == "ensemble":
        code = ctx.ensembles.get(_cle(texte))
        if code:
            return code, None, None
        nouveau = code_ensemble(texte)
        return nouveau, None, {"type": "ensemble", "code": nouveau, "nom": texte}
    code = ctx.listes.get(champ, {}).get(_cle(texte))
    if code:
        return code, None, None
    nouveau = listes.code_depuis_libelle(texte)
    return (
        nouveau,
        None,
        {"type": "valeur_liste", "liste": champ, "code": nouveau, "libelle": texte},
    )


def _attribut(champ: str, valeur: Any, texte: str, ctx: Contexte) -> tuple[Any, str | None]:
    """Valeur d'un attribut dans sa forme stockée : nombre, code de liste, '1' ou '0', texte."""
    attribut = ctx.attributs[champ.removeprefix(attributs.PREFIXE_CHAMP)]
    nom = attribut["libelle"]
    if attribut["type"] == "nombre":
        nombre = attributs.lire_nombre(valeur if isinstance(valeur, int | float) else texte)
        return (
            (nombre, None)
            if nombre is not None
            else (None, f"{nom} : « {texte} » n'est pas un nombre")
        )
    if attribut["type"] == "booleen":
        booleen = attributs.lire_booleen(texte)
        return (booleen, None) if booleen else (None, f"{nom} : « {texte} » ne vaut ni oui ni non")
    if attribut["type"] == "liste":
        for possible in attribut["valeurs"]:
            if _cle(texte) in (_cle(possible["code"]), _cle(possible["libelle"])):
                return possible["code"], None
        return None, f"{nom} : « {texte} » n'est pas une valeur de la liste"
    return texte, None


def _convertir(champ: str, valeur: Any, ctx: Contexte) -> tuple[Any, str | None, dict | None]:
    """Valeur typée d'une cellule, erreur en français, entité à créer éventuelle."""
    texte = texte_cellule(valeur)
    if texte is None:
        return None, None, None
    if champ.startswith(attributs.PREFIXE_CHAMP):
        return (*_attribut(champ, valeur, texte, ctx), None)
    nature, libelle = NATURES[champ], LIBELLES[champ]
    if nature in ("texte", "id"):
        return texte, None, None
    if nature in ("bloc", "fournisseur", "liste", "ensemble"):
        return _reference(champ, texte, ctx)
    if nature == "base":
        base = texte.upper()
        if base in ("HT", "TTC"):
            return base, None, None
        return None, f"{libelle} : « {texte} » (HT ou TTC attendu)", None
    est_nombre = isinstance(valeur, int | float) and not isinstance(valeur, bool)
    nombre = valeur if est_nombre else _nombre(texte)
    if nombre is None or nombre < 0:
        return None, f"{libelle} : « {texte} » n'est pas un nombre positif", None
    if nature == "entier":
        if float(nombre).is_integer():
            return int(nombre), None, None
        return None, f"{libelle} : « {texte} » n'est pas un nombre entier", None
    if nature == "taux":
        taux = nombre / 100 if nombre >= 1 else float(nombre)
        return (taux, None, None) if taux < 1 else (None, f"{libelle} : « {texte} » invalide", None)
    return float(nombre), None, None


def convert_ligne(numero: int, brut: dict[str, Any], ctx: Contexte) -> LigneLue | None:
    """Convertit une ligne ; None si elle est entièrement vide."""
    textes = {champ: texte_cellule(v) for champ, v in brut.items()}
    if all(v is None for v in textes.values()):
        return None
    lue = LigneLue(numero, {k: v for k, v in textes.items() if v is not None})
    for champ, valeur in brut.items():
        converti, erreur, entite = _convertir(champ, valeur, ctx)
        if erreur:
            lue.erreurs.append(erreur)
        if entite:
            lue.entites.append(entite)
        lue.valeurs[champ] = converti
    return lue


def complete_nouveau(lue: LigneLue, ctx: Contexte) -> None:
    """Valeurs par défaut et champs obligatoires d'un composant à créer."""
    valeurs = lue.valeurs
    if valeurs.get("bloc_code") is None and ctx.bloc_defaut:
        valeurs["bloc_code"] = ctx.bloc_defaut
    if valeurs.get("bloc_code") is None and not any("bloc" in e.lower() for e in lue.erreurs):
        lue.erreurs.append("champ obligatoire manquant : Bloc")
    # Dans un modèle d'ensemble, la quantité de l'ensemble tient lieu de besoin.
    if valeurs.get("qte_besoin") is None and valeurs.get("qte_ensemble"):
        valeurs["qte_besoin"] = valeurs["qte_ensemble"]
    for champ, libelle in OBLIGATOIRES.items():
        deja_signale = any(champ in e or libelle in e for e in lue.erreurs)
        if valeurs.get(champ) is None and not deja_signale:
            lue.erreurs.append(f"champ obligatoire manquant : {libelle}")
    valeurs["statut_appro"] = valeurs.get("statut_appro") or "Non lance"
    valeurs["base_prix_releve"] = valeurs.get("base_prix_releve") or "HT"
    if valeurs.get("taux_tva") is None:
        valeurs["taux_tva"] = ctx.taux_defaut
    for champ in ("qte_rechange", "qte_disponible"):
        valeurs[champ] = valeurs.get(champ) or 0
