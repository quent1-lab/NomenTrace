"""Comptes, connexion, droits et protections HTTP (phase 13).

Chaque test démarre une application en mode connecté sur le jeu d'essai, avec une base des
comptes temporaire. Le coût de scrypt est abaissé pour que les tests restent rapides : les
paramètres sont écrits dans chaque empreinte, le code de vérification est le même.
"""

import sqlite3
from collections.abc import Iterator
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import __main__ as lancement
from backend import comptes as ligne_de_commande
from backend import config, db
from backend.main import MODULES, create_app
from backend.services import authentification, comptes, droits

ADRESSE = "https://nomentrace.test"
PROJET = "essai"
MOT_DE_PASSE = "un mot de passe assez long"


@pytest.fixture(autouse=True)
def scrypt_rapide(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(authentification, "SCRYPT_N", 2**10)


def navigateur(app: FastAPI, origine: str | None = ADRESSE) -> TestClient:
    """Un navigateur distinct (ses propres cookies) sur une page de l'instance."""
    entetes = {"Origin": origine} if origine else {}
    return TestClient(app, base_url=ADRESSE, headers=entetes)


@pytest.fixture
def chemin_comptes(tmp_path: Path) -> Path:
    return tmp_path / "comptes.db"


@pytest.fixture
def app(
    conn_essai: sqlite3.Connection, chemin_base: Path, dossier_echange: Path, chemin_comptes: Path
) -> Iterator[FastAPI]:
    conn_essai.close()
    application = create_app(
        chemin_base,
        dossier_echange=dossier_echange,
        mode_local=False,
        chemin_comptes=chemin_comptes,
        code_projet=PROJET,
    )
    with navigateur(application):  # démarrage : migrations des deux bases
        yield application


def creer_compte(
    app: FastAPI,
    identifiant: str,
    role: str,
    blocs: tuple[str, ...] = (),
    permissions: tuple[str, ...] = (),
    projet: str = PROJET,
) -> int:
    """Crée un compte et choisit son mot de passe par le lien d'invitation."""
    conn = db.connect(app.state.chemin_comptes)
    try:
        cree = comptes.create_utilisateur(
            conn,
            projet,
            {
                "identifiant": identifiant,
                "nom": identifiant.split("@")[0],
                "role": role,
                "blocs": list(blocs),
                "permissions": list(permissions),
            },
            "test",
        )
        return comptes.accept_invitation(conn, cree["jeton"], MOT_DE_PASSE)
    finally:
        conn.close()


def connecter(app: FastAPI, identifiant: str, mot_de_passe: str = MOT_DE_PASSE) -> TestClient:
    client = navigateur(app)
    reponse = client.post(
        "/api/session", json={"identifiant": identifiant, "mot_de_passe": mot_de_passe}
    )
    assert reponse.status_code == 200, reponse.text
    return client


@pytest.fixture
def admin(app: FastAPI) -> TestClient:
    creer_compte(app, "admin@ecole.test", "administrateur")
    return connecter(app, "admin@ecole.test")


@pytest.fixture
def lecteur(app: FastAPI) -> TestClient:
    creer_compte(app, "lecteur@ecole.test", "lecteur")
    return connecter(app, "lecteur@ecole.test")


@pytest.fixture
def contributeur(app: FastAPI) -> TestClient:
    """Contributeur du bloc IHM, sans permission supplémentaire."""
    creer_compte(app, "ihm@ecole.test", "contributeur", blocs=("IHM",))
    return connecter(app, "ihm@ecole.test")


# --- Mots de passe ------------------------------------------------------------------------------


def test_empreinte_scrypt_salee_et_verifiee() -> None:
    premiere = authentification.hash_mot_de_passe(MOT_DE_PASSE)
    seconde = authentification.hash_mot_de_passe(MOT_DE_PASSE)
    assert premiere.startswith("scrypt$1024$8$1$")
    assert premiere != seconde  # sel propre à chaque empreinte
    assert MOT_DE_PASSE not in premiere
    assert authentification.verify_mot_de_passe(MOT_DE_PASSE, premiere)
    assert not authentification.verify_mot_de_passe(MOT_DE_PASSE + "x", premiere)
    assert not authentification.verify_mot_de_passe(MOT_DE_PASSE, "illisible")


def test_parametres_owasp_par_defaut(monkeypatch: pytest.MonkeyPatch) -> None:
    """Les tests abaissent le coût ; le code livré garde N = 2^17, r = 8, p = 1."""
    monkeypatch.undo()
    parametres = (authentification.SCRYPT_N, authentification.SCRYPT_R, authentification.SCRYPT_P)
    assert parametres == (2**17, 8, 1)


# --- Connexion ----------------------------------------------------------------------------------


def test_connexion_ouvre_une_session_securisee(app: FastAPI) -> None:
    creer_compte(app, "admin@ecole.test", "administrateur")
    client = navigateur(app)
    reponse = client.post(
        "/api/session", json={"identifiant": "ADMIN@ecole.test ", "mot_de_passe": MOT_DE_PASSE}
    )
    assert reponse.status_code == 200
    cookie = reponse.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=lax" in cookie and "secure" in cookie
    session = client.get("/api/session").json()
    assert session["utilisateur"]["role"] == "administrateur"
    assert session["nom_projet"] == "Projet d'essai"
    # Le jeton n'est stocké que par son empreinte.
    jeton = client.cookies.get("nomentrace_session")
    conn = db.connect(app.state.chemin_comptes)
    stocke = conn.execute("SELECT COUNT(*) FROM session WHERE empreinte = ?", (jeton,))
    assert stocke.fetchone()[0] == 0
    conn.close()


def test_mot_de_passe_faux_ou_compte_inconnu_meme_refus(app: FastAPI) -> None:
    creer_compte(app, "admin@ecole.test", "administrateur")
    client = navigateur(app)
    faux = client.post(
        "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": "mauvais mot"}
    )
    inconnu = client.post(
        "/api/session", json={"identifiant": "personne@ecole.test", "mot_de_passe": "mauvais mot"}
    )
    assert faux.status_code == inconnu.status_code == 401
    assert faux.json() == inconnu.json() == {"erreur": authentification.MESSAGE_ECHEC}
    assert client.get("/api/composants").status_code == 401


def test_blocage_apres_cinq_echecs(app: FastAPI) -> None:
    creer_compte(app, "admin@ecole.test", "administrateur")
    client = navigateur(app)
    for _ in range(5):
        reponse = client.post(
            "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": "mauvais"}
        )
        assert reponse.status_code == 401
    # Même le bon mot de passe est refusé le temps du blocage.
    bloque = client.post(
        "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": MOT_DE_PASSE}
    )
    assert bloque.status_code == 429
    # Un identifiant inconnu se bloque de la même façon : le blocage ne trahit pas les comptes.
    for _ in range(5):
        client.post("/api/session", json={"identifiant": "x@ecole.test", "mot_de_passe": "m"})
    inconnu = client.post("/api/session", json={"identifiant": "x@ecole.test", "mot_de_passe": "m"})
    assert inconnu.status_code == 429


def test_blocage_par_adresse() -> None:
    limiteur = authentification.Limiteur(3)
    for _ in range(3):
        assert not limiteur.bloque("10.0.0.1")
        limiteur.noter_echec("10.0.0.1")
    assert limiteur.bloque("10.0.0.1")
    assert not limiteur.bloque("10.0.0.2")


def test_connexion_refusee_quand_toutes_les_places_sont_prises(
    app: FastAPI, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Une rafale de connexions ne doit pas bloquer l'application : les places en trop sont
    refusées tout de suite (503), sans occuper un fil de traitement.
    """
    creer_compte(app, "admin@ecole.test", "administrateur")
    # Toutes les places sont occupées : une connexion de plus est refusée aussitôt.
    pris = [
        authentification._places.acquire() for _ in range(authentification.CONNEXIONS_SIMULTANEES)
    ]
    try:
        reponse = navigateur(app).post(
            "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": MOT_DE_PASSE}
        )
        assert reponse.status_code == 503
    finally:
        for _ in pris:
            authentification._places.release()
    # Une fois les places libérées, la connexion repasse normalement.
    assert connecter(app, "admin@ecole.test").get("/api/composants").status_code == 200


def test_places_liberees_meme_en_cas_d_echec(app: FastAPI) -> None:
    """La place réservée est rendue quel que soit le résultat, sinon elles s'épuiseraient."""
    creer_compte(app, "admin@ecole.test", "administrateur")
    client = navigateur(app)
    for _ in range(authentification.CONNEXIONS_SIMULTANEES + 3):
        client.post(
            "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": "faux"}
        )
    libres = [
        authentification._places.acquire(blocking=False)
        for _ in range(authentification.CONNEXIONS_SIMULTANEES)
    ]
    assert all(libres), "des places n'ont pas été rendues après un échec de connexion"
    for ok in libres:
        if ok:
            authentification._places.release()


def test_deconnexion_ferme_la_session(app: FastAPI, admin: TestClient) -> None:
    jeton = admin.cookies.get("nomentrace_session")
    assert admin.delete("/api/session").status_code == 200
    rejoue = navigateur(app)
    rejoue.cookies.set("nomentrace_session", jeton)
    assert rejoue.get("/api/composants").status_code == 401


def test_compte_sans_acces_au_projet_refuse(app: FastAPI) -> None:
    creer_compte(app, "autre@ecole.test", "administrateur", projet="autre-projet")
    client = navigateur(app)
    reponse = client.post(
        "/api/session", json={"identifiant": "autre@ecole.test", "mot_de_passe": MOT_DE_PASSE}
    )
    assert reponse.status_code == 403
    assert "nomentrace_session" not in client.cookies


def test_session_d_un_compte_sans_acces_ne_lit_rien(app: FastAPI, admin: TestClient) -> None:
    """Un accès retiré en cours de session ferme aussitôt les données du projet."""
    id_ = creer_compte(app, "parti@ecole.test", "lecteur")
    client = connecter(app, "parti@ecole.test")
    conn = db.connect(app.state.chemin_comptes)
    conn.execute("DELETE FROM acces WHERE utilisateur_id = ?", (id_,))
    conn.close()
    assert client.get("/api/composants").status_code == 403
    assert client.get("/api/sante").json()["nom_projet"] is None


# --- Invitations --------------------------------------------------------------------------------


def _invitation(app: FastAPI, identifiant: str, role: str = "lecteur") -> tuple[int, str]:
    conn = db.connect(app.state.chemin_comptes)
    cree = comptes.create_utilisateur(
        conn, PROJET, {"identifiant": identifiant, "nom": identifiant, "role": role}, "test"
    )
    conn.close()
    return cree["utilisateur"]["id"], cree["jeton"]


def test_invitation_choisit_le_mot_de_passe_et_connecte(app: FastAPI) -> None:
    _, jeton = _invitation(app, "nouveau@ecole.test")
    client = navigateur(app)
    assert client.post("/api/invitation/verifier", json={"jeton": jeton}).json() == {
        "identifiant": "nouveau@ecole.test",
        "nom": "nouveau@ecole.test",
    }
    court = client.post("/api/invitation", json={"jeton": jeton, "mot_de_passe": "court"})
    assert court.status_code == 400
    reponse = client.post("/api/invitation", json={"jeton": jeton, "mot_de_passe": MOT_DE_PASSE})
    assert reponse.status_code == 200
    assert client.get("/api/composants").status_code == 200
    # Usage unique.
    rejoue = navigateur(app).post(
        "/api/invitation", json={"jeton": jeton, "mot_de_passe": MOT_DE_PASSE + "!"}
    )
    assert rejoue.status_code == 410


@pytest.mark.parametrize(
    ("mot_de_passe", "accepte"),
    [
        ("Court1!", False),  # trop court
        ("toutenminuscule", False),  # ni majuscule, ni chiffre, ni symbole, pas une phrase
        ("Robot-Chassis-7", True),  # composé
        ("le robot monte ses roues avant", True),  # phrase de passe
        ("Motdepasse2026!", False),  # mot de passe courant à peine décoré
        ("Aaaaaaaaaaa1!", False),  # trop peu de caractères différents
        ("Marie.Dupont-26", False),  # reprend le nom
        ("Mdupont#Robot42", False),  # reprend l'adresse mail
    ],
)
def test_exigences_du_mot_de_passe(mot_de_passe: str, accepte: bool) -> None:
    exigences = authentification.exigences_mot_de_passe(
        mot_de_passe, "mdupont@ecole.fr", "Marie Dupont"
    )
    assert all(remplie for _, remplie in exigences) is accepte, exigences


def test_invitation_refuse_un_mot_de_passe_faible_en_citant_les_exigences(app: FastAPI) -> None:
    _, jeton = _invitation(app, "faible@ecole.test")
    reponse = navigateur(app).post(
        "/api/invitation", json={"jeton": jeton, "mot_de_passe": "Motdepasse2026!"}
    )
    assert reponse.status_code == 400
    assert "mot de passe courant" in reponse.json()["erreur"]
    # Le lien reste utilisable après un refus : rien n'a été consommé.
    ok = navigateur(app).post(
        "/api/invitation", json={"jeton": jeton, "mot_de_passe": MOT_DE_PASSE}
    )
    assert ok.status_code == 200


def test_invitation_expiree_refusee(app: FastAPI) -> None:
    _, jeton = _invitation(app, "tard@ecole.test")
    conn = db.connect(app.state.chemin_comptes)
    hier = (datetime.now() - timedelta(days=1)).isoformat(timespec="seconds")
    conn.execute("UPDATE invitation SET expire_le = ?", (hier,))
    conn.close()
    reponse = navigateur(app).post(
        "/api/invitation", json={"jeton": jeton, "mot_de_passe": MOT_DE_PASSE}
    )
    assert reponse.status_code == 410


def test_nouveau_lien_efface_le_mot_de_passe_et_les_sessions(
    app: FastAPI, admin: TestClient
) -> None:
    id_ = creer_compte(app, "oubli@ecole.test", "lecteur")
    session = connecter(app, "oubli@ecole.test")
    reponse = admin.post(f"/api/utilisateurs/{id_}/invitation")
    assert reponse.status_code == 200 and reponse.json()["jeton"]
    assert session.get("/api/composants").status_code == 401
    ancien = navigateur(app).post(
        "/api/session", json={"identifiant": "oubli@ecole.test", "mot_de_passe": MOT_DE_PASSE}
    )
    assert ancien.status_code == 401


# --- Gestion des comptes ------------------------------------------------------------------------


def test_admin_cree_et_desactive_un_compte(app: FastAPI, admin: TestClient) -> None:
    reponse = admin.post(
        "/api/utilisateurs",
        json={
            "identifiant": "Marie@Ecole.test",
            "nom": "Marie",
            "role": "contributeur",
            "blocs": ["IHM", "ALI"],
            "permissions": ["achats"],
        },
    )
    assert reponse.status_code == 201
    cree = reponse.json()
    assert cree["utilisateur"]["identifiant"] == "marie@ecole.test"
    assert cree["utilisateur"]["blocs"] == ["ALI", "IHM"]
    client = navigateur(app)
    client.post("/api/invitation", json={"jeton": cree["jeton"], "mot_de_passe": MOT_DE_PASSE})
    assert client.get("/api/composants").status_code == 200
    id_ = cree["utilisateur"]["id"]
    assert admin.patch(f"/api/utilisateurs/{id_}", json={"actif": False}).status_code == 200
    assert client.get("/api/composants").status_code == 401
    liste = admin.get("/api/utilisateurs").json()
    assert {u["nom"]: u["actif"] for u in liste}["Marie"] is False


def test_bloc_inconnu_ou_doublon_refuse(admin: TestClient) -> None:
    corps = {"identifiant": "a@ecole.test", "nom": "A", "role": "contributeur"}
    assert admin.post("/api/utilisateurs", json={**corps, "blocs": ["XYZ"]}).status_code == 400
    assert admin.post("/api/utilisateurs", json=corps).status_code == 201
    assert admin.post("/api/utilisateurs", json=corps).status_code == 409
    autre = {**corps, "identifiant": "b@ecole.test", "nom": "a"}
    assert admin.post("/api/utilisateurs", json=autre).status_code == 409  # nom distinctif


def test_dernier_administrateur_protege(app: FastAPI, admin: TestClient) -> None:
    moi = admin.get("/api/session").json()["utilisateur"]["id"]
    assert admin.patch(f"/api/utilisateurs/{moi}", json={"role": "lecteur"}).status_code == 409
    assert admin.patch(f"/api/utilisateurs/{moi}", json={"actif": False}).status_code == 409
    second = creer_compte(app, "second@ecole.test", "administrateur")
    assert admin.patch(f"/api/utilisateurs/{second}", json={"actif": False}).status_code == 200
    assert admin.patch(f"/api/utilisateurs/{moi}", json={"role": "lecteur"}).status_code == 409


# --- Droits par rôle ----------------------------------------------------------------------------


def test_toutes_les_routes_ont_une_regle() -> None:
    """Refus par défaut : aucune route ne doit dépendre de la règle implicite."""
    routes = {
        (methode, route.path)
        for module in MODULES
        for route in module.router.routes
        for methode in route.methods
        if methode != "HEAD"
    }
    assert routes - set(droits.REGLES) == set(), "routes sans règle d'accès"
    assert set(droits.REGLES) - routes == set(), "règles sans route"


def test_route_non_declaree_reservee_a_l_administrateur() -> None:
    assert droits.get_regle("POST", "/api/nouvelle-route") == droits.ADMIN


MOUVEMENT = {
    "date": "2026-02-01",
    "composant_id": "ESSAI-ALI-001",
    "type_mouvement": "Inventaire",
    "sens": "Entree",
    "qte": 1,
}

# (méthode, chemin, corps, statut du lecteur, du contributeur IHM, de l'administrateur)
CAS: tuple[tuple[str, str, dict | None, int, int, int], ...] = (
    ("GET", "/api/composants", None, 200, 200, 200),
    ("GET", "/api/pilotage", None, 200, 200, 200),
    ("GET", "/api/export/classeur", None, 200, 200, 200),
    ("GET", "/api/historique", None, 200, 200, 200),
    ("PATCH", "/api/composants/ESSAI-IHM-001", {"note_technique": "vu"}, 403, 200, 200),
    ("PATCH", "/api/composants/ESSAI-ALI-001", {"note_technique": "vu"}, 403, 403, 200),
    ("PUT", "/api/composants/ESSAI-ALI-001/attributs", {"valeurs": {}}, 403, 403, 200),
    ("POST", "/api/composants/ESSAI-IHM-001/reclasser", {"bloc_code": "ALI"}, 403, 403, 201),
    ("POST", "/api/mouvements", MOUVEMENT, 403, 201, 201),
    ("POST", "/api/commandes", {"fournisseur_nom": "Mouser"}, 403, 403, 201),
    ("POST", "/api/ensembles", {"code": "CHASSIS", "nom": "Châssis"}, 403, 403, 201),
    ("POST", "/api/fournisseurs", {"nom": "Nouveau"}, 403, 201, 201),
    ("PATCH", "/api/fournisseurs/Mouser", {"pays": "France"}, 403, 403, 200),
    ("POST", "/api/blocs", {"code": "NEW", "nom": "Nouveau"}, 403, 403, 201),
    ("POST", "/api/listes/criticite", {"libelle": "Majeure"}, 403, 403, 201),
    ("PATCH", "/api/parametres", {"nom_projet": "Renommé"}, 403, 403, 200),
    ("GET", "/api/sauvegardes", None, 403, 403, 200),
    ("GET", "/api/sauvegardes/archive", None, 403, 403, 200),
    ("POST", "/api/nettoyage/journal/purge", {"confirmer": True}, 403, 403, 200),
    ("GET", "/api/utilisateurs", None, 403, 403, 200),
)


@pytest.mark.parametrize(("methode", "chemin", "corps", "l", "c", "a"), CAS)
def test_droits_par_role(
    app: FastAPI,
    methode: str,
    chemin: str,
    corps: dict | None,
    l: int,  # noqa: E741
    c: int,
    a: int,
) -> None:
    creer_compte(app, "admin@ecole.test", "administrateur")
    creer_compte(app, "lecteur@ecole.test", "lecteur")
    creer_compte(app, "ihm@ecole.test", "contributeur", blocs=("IHM",))
    for identifiant, attendu in (
        ("lecteur@ecole.test", l),
        ("ihm@ecole.test", c),
        ("admin@ecole.test", a),
    ):
        client = connecter(app, identifiant)
        if corps and "nom" in corps and methode == "POST":
            corps = {**corps, "nom": f"{corps['nom']} {identifiant}"}  # une création par rôle
        reponse = client.request(methode, chemin, json=corps)
        assert reponse.status_code == attendu, (identifiant, reponse.text)


def test_non_connecte_ne_lit_que_la_sante(app: FastAPI) -> None:
    client = navigateur(app)
    assert client.get("/api/composants").status_code == 401
    assert client.get("/api/documents/1/fichier").status_code == 401
    sante = client.get("/api/sante").json()
    assert sante["statut"] == "ok" and sante["nom_projet"] is None
    # L'interface elle-même reste servie : c'est elle qui affiche la page de connexion.
    assert client.get("/connexion.html").status_code == 200
    assert client.get("/docs").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_contributeur_cree_seulement_dans_ses_blocs(contributeur: TestClient) -> None:
    composant = {
        "bloc_code": "ALI",
        "fonction": "Essai",
        "designation": "Pièce",
        "mode_appro": "Achat",
        "qte_besoin": 1,
    }
    assert contributeur.post("/api/composants", json=composant).status_code == 403
    cree = contributeur.post("/api/composants", json={**composant, "bloc_code": "IHM"})
    assert cree.status_code == 201


def test_permission_achats(app: FastAPI) -> None:
    creer_compte(app, "achats@ecole.test", "contributeur", blocs=("IHM",), permissions=("achats",))
    client = connecter(app, "achats@ecole.test")
    commande = client.post("/api/commandes", json={"fournisseur_nom": "Mouser"})
    assert commande.status_code == 201
    numero = commande.json()["numero"]
    # Une ligne peut porter un composant d'un autre bloc : la commande est collective.
    ligne = client.post(
        f"/api/commandes/{numero}/lignes",
        json={"composant_id": "ESSAI-ALI-001", "qte_commandee": 1, "pu_ht_devis": 2.5},
    )
    assert ligne.status_code == 201, ligne.text
    # Un fournisseur proposé reste à valider, et la validation revient à l'administrateur.
    cree = client.post("/api/fournisseurs", json={"nom": "Proposé", "statut": "Valide"})
    assert cree.json()["statut"] == "A valider"
    assert client.patch("/api/fournisseurs/Proposé", json={"statut": "Valide"}).status_code == 403
    assert client.patch("/api/fournisseurs/Proposé", json={"pays": "France"}).status_code == 200
    assert client.delete("/api/fournisseurs/Proposé").status_code == 403


def test_permission_ensembles_sans_budget(app: FastAPI, admin: TestClient) -> None:
    creer_compte(app, "ens@ecole.test", "contributeur", blocs=("IHM",), permissions=("ensembles",))
    client = connecter(app, "ens@ecole.test")
    assert client.post("/api/ensembles", json={"code": "BRAS", "nom": "Bras"}).status_code == 201
    budget = {"code": "TETE", "nom": "Tête", "budget_cible_ht": 100}
    assert client.post("/api/ensembles", json=budget).status_code == 403
    assert client.patch("/api/ensembles/BRAS", json={"nom": "Bras droit"}).status_code == 200
    assert client.patch("/api/ensembles/BRAS", json={"budget_cible_ht": 50}).status_code == 403
    assert admin.patch("/api/ensembles/BRAS", json={"budget_cible_ht": 50}).status_code == 200
    # Affectations : le composant doit être dans ses blocs, quel que soit l'ensemble.
    ihm = {"composant_id": "ESSAI-IHM-001", "qte": 1}
    ali = {"composant_id": "ESSAI-ALI-001", "qte": 1}
    assert client.post("/api/ensembles/BRAS/affectations", json=ihm).status_code == 201
    assert client.post("/api/ensembles/BRAS/affectations", json=ali).status_code == 403
    id_ali = admin.post("/api/ensembles/BRAS/affectations", json=ali).json()["affectation_id"]
    assert client.patch(f"/api/affectations/{id_ali}", json={"qte": 2}).status_code == 403
    assert client.delete(f"/api/affectations/{id_ali}").status_code == 403
    assert client.delete("/api/affectations/abc").status_code == 404


def test_import_limite_aux_blocs_du_contributeur(app: FastAPI, contributeur: TestClient) -> None:
    conn = db.connect(app.state.chemin_base)
    lot = db.insert_row(
        conn,
        "import_lot",
        {"depot": 0, "nom_fichier": "equipe.xlsx", "utilisateur": "ihm", "statut": "analyse"},
    )
    conn.execute("UPDATE import_lot SET depot = id WHERE id = ?", (lot,))
    lignes = {}
    for composant in ("ESSAI-IHM-001", "ESSAI-ALI-001"):
        actuel = conn.execute(
            "SELECT note_technique FROM composant WHERE id = ?", (composant,)
        ).fetchone()[0]
        donnees = (
            '{"fichier": "equipe.xlsx", "differences": {"note_technique": {"actuel": '
            + ("null" if actuel is None else f'"{actuel}"')
            + ', "propose": "revu"}}}'
        )
        lignes[composant] = db.insert_row(
            conn,
            "import_ligne",
            {
                "lot_id": lot,
                "numero_ligne": len(lignes) + 2,
                "categorie": "MODIFIE",
                "composant_id": composant,
                "donnees_json": donnees,
            },
        )
    conn.close()
    decisions = {str(i): {"champs": ["note_technique"]} for i in lignes.values()}
    reponse = contributeur.post(f"/api/imports/{lot}/appliquer", json={"decisions": decisions})
    assert reponse.status_code == 200, reponse.text
    resultat = reponse.json()
    assert resultat["appliquees"] == 1
    assert len(resultat["refus"]) == 1 and "ALI" in resultat["refus"][0]
    conn = db.connect(app.state.chemin_base)
    notes = dict(
        conn.execute(
            "SELECT id, note_technique FROM composant WHERE id IN (?, ?)", tuple(lignes)
        ).fetchall()
    )
    conn.close()
    assert notes["ESSAI-IHM-001"] == "revu"
    assert notes["ESSAI-ALI-001"] != "revu"


# --- Traçabilité et accès concurrents -----------------------------------------------------------


def test_journal_et_historique_portent_l_utilisateur(contributeur: TestClient) -> None:
    contributeur.patch("/api/composants/ESSAI-IHM-001", json={"note_technique": "par ihm"})
    derniere = contributeur.get("/api/historique").json()["lignes"][0]
    assert derniere["utilisateur"] == "ihm"
    export = contributeur.get("/api/historique/export")
    feuille = load_workbook(BytesIO(export.content)).active
    entetes = [c.value for c in feuille[1]]
    assert feuille.cell(2, entetes.index("Utilisateur") + 1).value == "ihm"


def test_modification_concurrente_signalee(app: FastAPI, admin: TestClient) -> None:
    creer_compte(app, "ihm@ecole.test", "contributeur", blocs=("IHM",))
    autre = connecter(app, "ihm@ecole.test")
    lu = admin.get("/api/composants/ESSAI-IHM-001").json()["composant"]["modifie_le"]
    conn = db.connect(app.state.chemin_base)
    conn.execute(
        "UPDATE composant SET modifie_le = '2026-02-02T08:00:00' WHERE id = 'ESSAI-IHM-001'"
    )
    conn.close()
    autre.patch(
        "/api/composants/ESSAI-IHM-001",
        json={"note_technique": "d'abord", "modifie_le": "2026-02-02T08:00:00"},
    )
    reponse = admin.patch(
        "/api/composants/ESSAI-IHM-001", json={"note_technique": "ensuite", "modifie_le": lu}
    )
    assert reponse.status_code == 409
    assert "ihm" in reponse.json()["erreur"]
    fiche = admin.get("/api/composants/ESSAI-IHM-001").json()
    assert fiche["composant"]["note_technique"] == "d'abord"


def test_connexion_attend_un_verrou_d_ecriture(tmp_path: Path) -> None:
    conn = db.connect(tmp_path / "a.db")
    assert conn.execute("PRAGMA busy_timeout").fetchone()[0] == 5000
    conn.close()


# --- Protections HTTP ---------------------------------------------------------------------------


def test_ecriture_d_une_autre_origine_refusee(app: FastAPI, admin: TestClient) -> None:
    jeton = admin.cookies.get("nomentrace_session")
    for origine in ("https://site-piege.test", None):
        piege = navigateur(app, origine)
        piege.cookies.set("nomentrace_session", jeton)
        reponse = piege.patch("/api/parametres", json={"nom_projet": "Piraté"})
        assert reponse.status_code == 403
    assert admin.get("/api/parametres").json()["nom_projet"] == "Projet d'essai"
    # La connexion elle-même exige l'origine : pas de connexion forcée depuis un autre site.
    reponse = navigateur(app, "https://site-piege.test").post(
        "/api/session", json={"identifiant": "admin@ecole.test", "mot_de_passe": MOT_DE_PASSE}
    )
    assert reponse.status_code == 403


def test_en_tetes_de_securite(admin: TestClient) -> None:
    for chemin in ("/", "/api/composants"):
        reponse = admin.get(chemin)
        politique = reponse.headers["content-security-policy"]
        assert "script-src 'self'" in politique and "frame-ancestors 'none'" in politique
        assert reponse.headers["x-content-type-options"] == "nosniff"
        assert reponse.headers["x-frame-options"] == "DENY"
    assert admin.get("/api/composants").headers["cache-control"] == "no-store"
    assert "max-age" in admin.get("/").headers["strict-transport-security"]


def test_lien_javascript_refuse(admin: TestClient) -> None:
    for lien in ("javascript:alert(1)", " JavaScript:alert(1)", "data:text/html,x"):
        reponse = admin.patch("/api/composants/ESSAI-IHM-001", json={"lien_produit": lien})
        assert reponse.status_code == 422
    for lien in ("https://exemple.test/p", "www.exemple.test"):
        reponse = admin.patch("/api/composants/ESSAI-IHM-001", json={"lien_produit": lien})
        assert reponse.status_code == 200


def test_export_excel_sans_formule(admin: TestClient) -> None:
    admin.patch("/api/composants/ESSAI-IHM-001", json={"designation": '=HYPERLINK("x")'})
    for chemin in ("/api/export/classeur", "/api/composants/export", "/api/historique/export"):
        classeur = load_workbook(BytesIO(admin.get(chemin).content))
        types = {c.data_type for f in classeur.worksheets for ligne in f.iter_rows() for c in ligne}
        assert "f" not in types, chemin


# --- Mode local ---------------------------------------------------------------------------------


def test_mode_local_refuse_hors_boucle_locale() -> None:
    assert lancement.check_configuration("127.0.0.1", True) is None
    assert lancement.check_configuration("0.0.0.0", False) is None
    refus = lancement.check_configuration("0.0.0.0", True)
    assert refus and "mode local" in refus


def test_mode_local_refuse_autre_poste_et_autre_nom(
    chemin_base: Path, dossier_echange: Path
) -> None:
    app = create_app(chemin_base, dossier_echange=dossier_echange, mode_local=True)
    with TestClient(app, base_url="http://127.0.0.1:8000", client=("192.168.1.20", 5000)) as poste:
        assert poste.get("/api/composants").status_code == 403
    # Nom de domaine détourné vers 127.0.0.1 (DNS rebinding) : l'hôte n'est pas local.
    piege = TestClient(app, base_url="http://piege.test:8000", client=("127.0.0.1", 5000))
    assert piege.get("/api/composants").status_code == 403
    local = TestClient(app, base_url="http://localhost:8000", client=("127.0.0.1", 5000))
    assert local.get("/api/session").json()["utilisateur"]["nom"] == "local"


def test_mode_local_journalise_local(client: TestClient) -> None:
    client.post("/api/blocs", json={"code": "NEW", "nom": "Nouveau"})
    lignes = client.get("/api/journal").json()
    assert lignes[0]["utilisateur"] == "local"


def test_mode_connecte_par_defaut(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(config, "MODE_LOCAL", False)
    app = create_app(tmp_path / "p.db", tmp_path / "e", chemin_comptes=tmp_path / "c.db")
    assert app.state.mode_local is False


# --- Ligne de commande --------------------------------------------------------------------------


def test_creer_admin_en_ligne_de_commande(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(config, "CHEMIN_COMPTES", tmp_path / "comptes.db")
    monkeypatch.setattr(config, "CODE_PROJET", PROJET)
    monkeypatch.setattr(config, "URL_PUBLIQUE", "https://nomentrace.test/")
    assert ligne_de_commande.main(["creer-admin", "Chef@Ecole.test", "--nom", "Chef"]) == 0
    sortie = capsys.readouterr().out
    assert "https://nomentrace.test/connexion.html#invitation=" in sortie
    # Relancée, la commande rétablit le compte et donne un lien neuf (secours).
    assert ligne_de_commande.main(["creer-admin", "chef@ecole.test"]) == 0
    conn = db.connect(config.CHEMIN_COMPTES)
    assert comptes.count_admins(conn, PROJET) == 1
    ouvertes = conn.execute("SELECT COUNT(*) FROM invitation WHERE utilisee_le IS NULL")
    assert ouvertes.fetchone()[0] == 1
    conn.close()
