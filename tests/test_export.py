"""Tests de l'export Excel et des sauvegardes."""

import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.services import export_excel, sauvegardes


def test_export_lisible_par_openpyxl(client_essai: TestClient, dossier_echange: Path) -> None:
    assert client_essai.post("/api/export").json() == {"statut": "ok"}
    classeur = load_workbook(dossier_echange / "exports" / "nomenclature.xlsx", read_only=True)
    assert classeur.sheetnames[:4] == ["Pilotage", "Blocs", "Ensembles", "Affectations"]
    assert "composant" in classeur.sheetnames
    pilotage = classeur["Pilotage"]
    assert "généré automatiquement" in pilotage["A1"].value
    valeurs = {ligne[0]: ligne[1] for ligne in pilotage.iter_rows(min_row=4, values_only=True)}
    assert valeurs["cout_ht"] == 3184.98
    assert classeur["composant"].max_row == 61
    classeur.close()


def test_export_verrouille_passe_en_attente(
    client_essai: TestClient, dossier_echange: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def remplacement_refuse(_source: str, _cible: object) -> None:
        raise PermissionError("fichier ouvert dans Excel")

    monkeypatch.setattr(export_excel.os, "replace", remplacement_refuse)
    assert client_essai.post("/api/export").json() == {"statut": "en_attente"}
    assert client_essai.get("/api/sante").json()["export_en_attente"] is True
    assert list((dossier_echange / "exports").glob("*.tmp")) == []
    monkeypatch.undo()
    assert client_essai.post("/api/export").json() == {"statut": "ok"}
    assert client_essai.get("/api/sante").json()["export_en_attente"] is False


def test_anti_rafale_un_seul_export(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    appels: list[float] = []
    monkeypatch.setattr(export_excel, "write_export", lambda *_: appels.append(time.monotonic()))
    planificateur = export_excel.PlanificateurExport(tmp_path / "b.db", tmp_path, delai=0.3)
    planificateur.start()
    try:
        debut = time.monotonic()
        for _ in range(5):
            planificateur.signaler()
            time.sleep(0.05)
        time.sleep(0.8)
    finally:
        planificateur.stop()
    assert len(appels) == 1
    assert appels[0] - debut >= 0.3


def test_sauvegarde_au_demarrage_et_rotation(
    client_essai: TestClient, chemin_base: Path, tmp_path: Path
) -> None:
    dossier = tmp_path / "sauvegardes_test"
    for i in range(25):
        (dossier / f"nomentrace_20260101_0000{i:02d}.db").parent.mkdir(exist_ok=True)
        (dossier / f"nomentrace_20260101_0000{i:02d}.db").write_bytes(b"")
    cible = sauvegardes.create_sauvegarde(chemin_base, dossier)
    assert cible is not None
    restantes = sauvegardes.list_sauvegardes(dossier)
    assert len(restantes) == 20
    assert restantes[0] == cible
