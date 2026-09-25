"""Import des fichiers de l'équipe : classement, doublons, application, traçabilité."""

import io
import sqlite3
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend import db

BLOCS = ("TR", "IHM", "ODB", "AGI", "OBS", "ALI", "BMP", "BUS", "MEC")


def _modele(client: TestClient, chemin: str) -> bytes:
    reponse = client.get(chemin)
    assert reponse.status_code == 200, reponse.text
    return reponse.content


def _modifier(contenu: bytes, identifiant: str, entete: str, valeur: Any) -> bytes:
    """Change une cellule de la ligne d'un composant dans un modèle."""
    classeur = load_workbook(io.BytesIO(contenu))
    feuille = classeur["Composants"]
    entetes = [c.value for c in feuille[1]]
    for ligne in feuille.iter_rows(min_row=2):
        if ligne[0].value == identifiant:
            ligne[entetes.index(entete)].value = valeur
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def _fichier_libre(lignes: list[dict[str, Any]]) -> bytes:
    """Fichier saisi à la main par l'équipe, sans feuille d'identification."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Composants"
    entetes = [
        "ID",
        "Bloc",
        "Fonction",
        "Désignation",
        "Réf fabricant",
        "Mode appro",
        "Qté besoin",
        "Fournisseur",
        "Criticité",
        "Ensemble",
        "Qté dans cet ensemble",
    ]
    feuille.append(entetes)
    for ligne in lignes:
        feuille.append([ligne.get(e) for e in entetes])
    sortie = io.BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


def _deposer(client: TestClient, fichiers: list[tuple[str, bytes]]) -> dict:
    reponse = client.post(
        "/api/imports",
        files=[("fichiers", (nom, contenu)) for nom, contenu in fichiers],
        data={"depose_par": "Test"},
    )
    assert reponse.status_code == 201, reponse.text
    return client.get(f"/api/imports/{reponse.json()['depot']}").json()


def _lignes(depot: dict, categorie: str) -> list[dict]:
    return [ligne for ligne in depot["lignes"] if ligne["categorie"] == categorie]


def _decisions_par_defaut(depot: dict) -> dict:
    """Ce que l'écran de revue coche d'office : tout accepter, action proposée."""
    decisions = {}
    for ligne in depot["lignes"]:
        d = ligne["donnees"]
        if ligne["categorie"] == "MODIFIE":
            decisions[str(ligne["id"])] = {"champs": list(d["differences"])}
        elif ligne["categorie"] in ("NOUVEAU", "DOUBLON"):
            decision = {"action": d["action_defaut"]}
            if d["action_defaut"] == "fusionner":
                meilleur = next(c for c in d["candidats"] if c["fusion"])
                decision["cible"] = (
                    {"composant": meilleur["id"]}
                    if meilleur["type"] == "composant"
                    else {"ligne": meilleur["ligne_id"]}
                )
            decisions[str(ligne["id"])] = decision
        elif ligne["categorie"].endswith("_AFFECTATION"):
            decisions[str(ligne["id"])] = {"action": "appliquer"}
    return decisions


def _appliquer(client: TestClient, depot: dict, decisions: dict) -> dict:
    reponse = client.post(f"/api/imports/{depot['depot']}/appliquer", json={"decisions": decisions})
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def _instantane(chemin_base: Path) -> tuple:
    conn = db.connect(chemin_base)
    try:
        return tuple(
            db.fetch_all(conn, f"SELECT * FROM {table} ORDER BY 1")  # noqa: S608
            for table in ("composant", "affectation", "ensemble")
        )
    finally:
        conn.close()


# --- Modèles et aller-retour ---------------------------------------------------------------


def test_neuf_modeles_redeposes_sans_modification(client_essai: TestClient) -> None:
    fichiers = [(f"{b}.xlsx", _modele(client_essai, f"/api/blocs/{b}/modele")) for b in BLOCS]
    depot = _deposer(client_essai, fichiers)
    assert depot["resume"] == {"IDENTIQUE": 60}
    assert len(depot["lots"]) == 9
    assert {lot["bloc_devine"] for lot in depot["lots"]} == set(BLOCS)


def test_modele_reconnu_quel_que_soit_son_nom(client_essai: TestClient) -> None:
    depot = _deposer(
        client_essai, [("renomme.xlsx", _modele(client_essai, "/api/blocs/ALI/modele"))]
    )
    assert depot["lots"][0]["bloc_devine"] == "ALI"


def test_prix_modifie_sur_le_bon_champ(client_essai: TestClient) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", 21.5
    )
    depot = _deposer(client_essai, [("obs.xlsx", contenu)])
    assert depot["resume"] == {"MODIFIE": 1, "IDENTIQUE": 3}
    modifie = _lignes(depot, "MODIFIE")[0]
    assert modifie["composant_id"] == "ESSAI-OBS-002"
    assert modifie["donnees"]["differences"] == {"pu_releve": {"actuel": 19.9, "propose": 21.5}}


def test_meme_prix_ecrit_autrement_n_est_pas_modifie(client_essai: TestClient) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", "19,90"
    )
    contenu = _modifier(contenu, "ESSAI-OBS-003", "PU relevé", 17.7500000001)
    depot = _deposer(client_essai, [("obs.xlsx", contenu)])
    assert depot["resume"] == {"IDENTIQUE": 4}


# --- Doublons -----------------------------------------------------------------------------


def test_reference_identique_certitude_100(client_essai: TestClient) -> None:
    ligne = {
        "Bloc": "TR",
        "Fonction": "Bus CAN",
        "Désignation": "Transceiver",
        "Réf fabricant": "MCP2562-E-P",
        "Mode appro": "Achat",
        "Qté besoin": 2,
    }
    depot = _deposer(client_essai, [("equipe.xlsx", _fichier_libre([ligne]))])
    doublon = _lignes(depot, "DOUBLON")[0]
    candidat = doublon["donnees"]["candidats"][0]
    assert (candidat["id"], candidat["niveau"], candidat["score"]) == ("ESSAI-TR-001", 1, 1.0)
    assert doublon["donnees"]["action_defaut"] == "fusionner"


def test_reference_citee_dans_le_libelle(client_essai: TestClient) -> None:
    ligne = {
        "Bloc": "AGI",
        "Fonction": "Communication",
        "Désignation": "Transceiver CAN MCP 2562",
        "Mode appro": "Achat",
        "Qté besoin": 1,
    }
    depot = _deposer(client_essai, [("agir.xlsx", _fichier_libre([ligne]))])
    candidat = _lignes(depot, "DOUBLON")[0]["donnees"]["candidats"][0]
    assert (candidat["id"], candidat["niveau"]) == ("ESSAI-TR-001", 2)


def test_garde_numerique_pas_de_fusion(client_essai: TestClient) -> None:
    ligne = {
        "Bloc": "ALI",
        "Fonction": "Conversion",
        "Désignation": "Convertisseur DC-DC 36 V",
        "Mode appro": "Achat",
        "Qté besoin": 1,
    }
    depot = _deposer(client_essai, [("ali.xlsx", _fichier_libre([ligne]))])
    doublon = _lignes(depot, "DOUBLON")[0]["donnees"]
    assert doublon["action_defaut"] == "creer"
    proches = {c["id"] for c in doublon["candidats"]}
    assert {"ESSAI-ALI-004", "ESSAI-ALI-005"} <= proches
    assert not any(c["fusion"] for c in doublon["candidats"])


def test_meme_nouveau_composant_dans_deux_fichiers(client_essai: TestClient) -> None:
    ligne = {
        "Bloc": "BUS",
        "Fonction": "Câblage capteurs",
        "Désignation": "Nappe blindée 8 conducteurs",
        "Réf fabricant": "NB-8C-2024",
        "Mode appro": "Achat",
        "Qté besoin": 2,
    }
    depot = _deposer(
        client_essai, [("a.xlsx", _fichier_libre([ligne])), ("b.xlsx", _fichier_libre([ligne]))]
    )
    assert depot["resume"] == {"NOUVEAU": 1, "DOUBLON": 1}
    candidat = _lignes(depot, "DOUBLON")[0]["donnees"]["candidats"][0]
    assert candidat["type"] == "ligne"
    resultat = _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    assert len(resultat["crees"]) == 1
    cree = client_essai.get(f"/api/composants/{resultat['crees'][0]}").json()["composant"]
    assert cree["qte_besoin"] == 4


# --- Lignes inexploitables ---------------------------------------------------------------


def test_bloc_inexistant_invalide_sans_bloquer_le_reste(client_essai: TestClient) -> None:
    lignes = [
        {
            "Bloc": "ZZZ",
            "Fonction": "X",
            "Désignation": "Pièce inconnue",
            "Mode appro": "Achat",
            "Qté besoin": 1,
        },
        {
            "Bloc": "MEC",
            "Fonction": "Fixation",
            "Désignation": "Vis CHC M4x12 inox",
            "Mode appro": "Achat",
            "Qté besoin": 40,
        },
        {"ID": "ESSAI-XXX-999", "Désignation": "Inconnu"},
    ]
    depot = _deposer(client_essai, [("mec.xlsx", _fichier_libre(lignes))])
    # L'ID inventé n'est pas rejeté comme inconnu : la ligne est traitée comme un nouveau
    # composant, invalide ici faute de fonction, de mode d'appro et de besoin.
    assert depot["resume"] == {"NOUVEAU": 1, "INVALIDE": 2}
    invalide = _lignes(depot, "INVALIDE")[0]
    assert "ZZZ" in invalide["donnees"]["raisons"][0]


def test_fichier_illisible(client_essai: TestClient) -> None:
    depot = _deposer(client_essai, [("faux.xlsx", b"ceci n'est pas un classeur")])
    assert depot["resume"] == {"INVALIDE": 1}


# --- Modèle d'ensemble ------------------------------------------------------------------------


def test_modele_d_ensemble_propose_les_affectations(client_essai: TestClient) -> None:
    client_essai.post("/api/ensembles", json={"code": "NACELLE", "nom": "Nacelle"})
    for composant, qte in (("ESSAI-TR-001", 2), ("ESSAI-ALI-003", 1)):
        corps = {"composant_id": composant, "qte": qte}
        client_essai.post("/api/ensembles/NACELLE/affectations", json=corps)
    modele = _modele(client_essai, "/api/ensembles/NACELLE/modele")
    contenu = _modifier(modele, "ESSAI-TR-001", "Qté dans cet ensemble", 3)
    contenu = _modifier(contenu, "ESSAI-ALI-003", "Qté dans cet ensemble", 0)
    depot = _deposer(client_essai, [("nacelle.xlsx", contenu)])
    assert depot["lots"][0]["ensemble_devine"] == "NACELLE"
    assert depot["resume"] == {"MODIF_AFFECTATION": 1, "SUPPRESSION_AFFECTATION": 1, "IDENTIQUE": 2}
    _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    lignes = client_essai.get("/api/ensembles/NACELLE/composants").json()
    assert [(ligne["composant_id"], ligne["qte_affectee"]) for ligne in lignes] == [
        ("ESSAI-TR-001", 3)
    ]


# --- Application ------------------------------------------------------------------------------


def test_tout_refuser_laisse_la_base_inchangee(client_essai: TestClient, chemin_base: Path) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", 25
    )
    nouveau = {
        "Bloc": "OBS",
        "Fonction": "Mesure",
        "Désignation": "Capteur de pression XP-9000",
        "Mode appro": "Achat",
        "Qté besoin": 1,
    }
    avant = _instantane(chemin_base)
    depot = _deposer(
        client_essai, [("obs.xlsx", contenu), ("libre.xlsx", _fichier_libre([nouveau]))]
    )
    assert _instantane(chemin_base) == avant
    refus = {str(ligne["id"]): {"champs": [], "action": "ignorer"} for ligne in depot["lignes"]}
    resultat = _appliquer(client_essai, depot, refus)
    assert resultat["appliquees"] == 0
    assert _instantane(chemin_base) == avant


def test_appliquer_puis_redeposer_donne_identique(client_essai: TestClient) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", 22
    )
    depot = _deposer(client_essai, [("obs_equipe.xlsx", contenu)])
    resultat = _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    assert resultat == {"appliquees": 1, "refus": [], "crees": []}
    assert _deposer(client_essai, [("obs_equipe.xlsx", contenu)])["resume"] == {"IDENTIQUE": 4}


def test_journal_retrouve_le_fichier(client_essai: TestClient) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", 23
    )
    depot = _deposer(client_essai, [("obs_de_quentin.xlsx", contenu)])
    _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    journal = client_essai.get("/api/composants/ESSAI-OBS-002").json()["historique"]
    entree = next(j for j in journal if j["champ"] == "pu_releve")
    assert (entree["origine"], entree["nom_fichier"], entree["nouvelle_valeur"]) == (
        "import",
        "obs_de_quentin.xlsx",
        "23.0",
    )


def test_modification_concurrente_refusee(client_essai: TestClient) -> None:
    contenu = _modifier(
        _modele(client_essai, "/api/blocs/OBS/modele"), "ESSAI-OBS-002", "PU relevé", 24
    )
    depot = _deposer(client_essai, [("obs.xlsx", contenu)])
    client_essai.patch("/api/composants/ESSAI-OBS-002", json={"pu_releve": 30})
    resultat = _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    assert resultat["appliquees"] == 0
    assert "modifié depuis l'analyse" in resultat["refus"][0]
    assert client_essai.get("/api/composants/ESSAI-OBS-002").json()["composant"]["pu_releve"] == 30


def test_depot_applique_une_seule_fois(client_essai: TestClient) -> None:
    depot = _deposer(client_essai, [("obs.xlsx", _modele(client_essai, "/api/blocs/OBS/modele"))])
    _appliquer(client_essai, depot, {})
    reponse = client_essai.post(f"/api/imports/{depot['depot']}/appliquer", json={"decisions": {}})
    assert reponse.status_code == 400


def test_sauvegarde_avant_application(client_essai: TestClient, dossier_echange: Path) -> None:
    avant = len(list((dossier_echange / "sauvegardes").glob("*.db")))
    depot = _deposer(client_essai, [("obs.xlsx", _modele(client_essai, "/api/blocs/OBS/modele"))])
    _appliquer(client_essai, depot, {})
    assert len(list((dossier_echange / "sauvegardes").glob("*.db"))) == avant + 1


def test_rien_n_est_ecrit_hors_tables_d_import(conn_essai: sqlite3.Connection) -> None:
    """L'analyse seule n'écrit aucune ligne de journal."""
    from backend.services import import_analyse

    avant = conn_essai.execute("SELECT COUNT(*) FROM journal").fetchone()[0]
    ligne = {
        "Bloc": "TR",
        "Fonction": "F",
        "Désignation": "D",
        "Mode appro": "Achat",
        "Qté besoin": 1,
    }
    import_analyse.analyse_depot(
        conn_essai,
        Path(conn_essai.execute("PRAGMA database_list").fetchone()[2]).parent,
        [("x.xlsx", _fichier_libre([ligne]))],
        None,
    )
    assert conn_essai.execute("SELECT COUNT(*) FROM journal").fetchone()[0] == avant


# --- Identifiant inventé ----------------------------------------------------------------------


def test_identifiant_invente_devient_un_nouveau_composant(client_essai: TestClient) -> None:
    ligne = {
        "ID": "ESSAI-OBS-099",
        "Bloc": "OBS",
        "Fonction": "Mesure",
        "Désignation": "Sonde PT100",
        "Mode appro": "Achat",
        "Qté besoin": 2,
    }
    depot = _deposer(client_essai, [("obs.xlsx", _fichier_libre([ligne]))])
    assert depot["resume"] == {"NOUVEAU": 1}
    nouveau = _lignes(depot, "NOUVEAU")[0]["donnees"]
    assert nouveau["id_indicatif"] == "ESSAI-OBS-005"
    assert "n'existe pas" in nouveau["avertissements"][0]


def test_identifiant_inconnu_sans_donnees_reste_inconnu(client_essai: TestClient) -> None:
    ligne = {"ID": "ESSAI-OBS-099", "Ensemble": "NACELLE", "Qté dans cet ensemble": 1}
    depot = _deposer(client_essai, [("obs.xlsx", _fichier_libre([ligne]))])
    assert depot["resume"]["INCONNU"] == 1


def test_modele_colonne_id_verrouillee(client_essai: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(_modele(client_essai, "/api/blocs/OBS/modele")))
    feuille = classeur["Composants"]
    assert feuille.protection.sheet
    assert feuille["A2"].protection.locked
    assert feuille["A40"].protection.locked
    assert not feuille["C40"].protection.locked
    entetes = [c.value for c in feuille[1]]
    assert entetes[-2:] == ["Ensemble", "Qté dans cet ensemble"]


# --- Nouvelles entités -------------------------------------------------------------------------


def _ligne_complete(**autres: Any) -> dict:
    return {
        "Bloc": "MEC",
        "Fonction": "Fixation",
        "Désignation": "Écrou nylstop M6",
        "Mode appro": "Achat",
        "Qté besoin": 20,
        **autres,
    }


def test_entites_inconnues_proposees_puis_creees(client_essai: TestClient) -> None:
    lignes = [
        _ligne_complete(
            **{
                "Fournisseur": "Bossard",
                "Criticité": "Critique",
                "Ensemble": "Châssis arrière",
                "Qté dans cet ensemble": 12,
            }
        ),
        _ligne_complete(
            **{
                "Désignation": "Rondelle M6",
                "Fournisseur": "bossard",
                "Ensemble": "Châssis arrière",
                "Qté dans cet ensemble": 24,
            }
        ),
    ]
    depot = _deposer(client_essai, [("mec.xlsx", _fichier_libre(lignes))])
    entites = [ligne["donnees"]["entite"] for ligne in _lignes(depot, "NOUVELLE_ENTITE")]
    assert entites == [
        {"type": "fournisseur", "nom": "Bossard"},
        {"type": "valeur_liste", "liste": "criticite", "code": "Critique", "libelle": "Critique"},
        {"type": "ensemble", "code": "CHASSIS-ARRIERE", "nom": "Châssis arrière"},
    ]
    decisions = _decisions_par_defaut(depot)
    decisions.update(
        {str(ligne["id"]): {"action": "creer"} for ligne in _lignes(depot, "NOUVELLE_ENTITE")}
    )
    resultat = _appliquer(client_essai, depot, decisions)
    assert resultat["refus"] == []
    assert len(resultat["crees"]) == 2
    bossard = [f for f in client_essai.get("/api/fournisseurs").json() if f["nom"] == "Bossard"]
    # Trouvé par l'équipe : créé « à valider ».
    assert [f["statut"] for f in bossard] == ["A valider"]
    ensemble = client_essai.get("/api/ensembles/CHASSIS-ARRIERE").json()
    assert (ensemble["nom"], ensemble["nb_pieces_total"]) == ("Châssis arrière", 36)
    criticites = [v["code"] for v in client_essai.get("/api/listes").json()["criticite"]]
    assert "Critique" in criticites


def test_entite_refusee_refuse_les_lignes_qui_la_citent(client_essai: TestClient) -> None:
    depot = _deposer(
        client_essai, [("mec.xlsx", _fichier_libre([_ligne_complete(Fournisseur="Inconnu SARL")]))]
    )
    decisions = _decisions_par_defaut(depot)
    entite = _lignes(depot, "NOUVELLE_ENTITE")[0]
    decisions[str(entite["id"])] = {"action": "ignorer"}
    resultat = _appliquer(client_essai, depot, decisions)
    assert resultat["crees"] == []
    assert "Inconnu SARL" in resultat["refus"][0]


def test_ensemble_existant_par_colonne(client_essai: TestClient) -> None:
    client_essai.post("/api/ensembles", json={"code": "MAT", "nom": "Mât"})
    ligne = {"ID": "ESSAI-TR-001", "Ensemble": "Mât", "Qté dans cet ensemble": 2}
    depot = _deposer(client_essai, [("tr.xlsx", _fichier_libre([ligne]))])
    creation = _lignes(depot, "CREATION_AFFECTATION")[0]
    assert (creation["donnees"]["ensemble"], creation["donnees"]["qte_proposee"]) == ("MAT", 2)
    assert "NOUVELLE_ENTITE" not in depot["resume"]
