"""Suivi des achats tenu par l'établissement : classeur imposé, en aller-retour.

Le modèle est construit ici, fictif, sur la forme d'un modèle d'établissement : en-tête,
légende des couleurs, titres de colonnes, lignes de données, ligne TOTAL avec ses sommes.
"""

import io
import json
from datetime import date, datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

TITRES = (
    "Fournisseur",
    "Ref matériel",
    "Désignation matériel",
    "Qté",
    "montant\n€  HT",
    "date devis",
    "date commande*",
    "date livraison prévue*",
    "date réception",
    "montant commande HT",
    "Commentaires",
)
COMMANDE = PatternFill("solid", fgColor="DAEEF3")
LIVRE = PatternFill("solid", fgColor="92D050")
BLEU = Font(color="FF0070C0")
NB_LIGNES = 4  # lignes de données prévues par le modèle fictif (10 à 13)


def _modele() -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "GROUPE"
    feuille["A1"], feuille["B1"] = "PROMO", "Promo d'essai"
    feuille["A2"], feuille["B2"] = "Groupe", "Équipe d'essai"
    feuille["E2"], feuille["F2"], feuille["G2"] = "devis", "commandé", "livré"
    feuille["F2"].fill, feuille["G2"].fill = COMMANDE, LIVRE
    feuille["A6"] = "SUIVI DES COMMANDES"
    for colonne, titre in enumerate(TITRES, start=1):
        feuille.cell(9, colonne, titre)
    trait = Side(style="thin")
    for rang in range(10, 10 + NB_LIGNES):
        for colonne in range(1, len(TITRES) + 1):
            feuille.cell(rang, colonne).border = Border(bottom=trait)
    total = 10 + NB_LIGNES
    feuille.cell(total, 4, "TOTAL")
    feuille.cell(total, 5, f"=SUM(E10:E{total - 1})")
    feuille.cell(total, 9, "TOTAL")
    feuille.cell(total, 10, f"=SUM(J10:J{total - 1})")
    feuille.cell(total + 3, 1, "Colonnes à droite : à compléter par l'établissement")
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _commande(client: TestClient, composants: list[str], transmise: bool = True) -> str:
    numero = client.post(
        "/api/commandes",
        json={
            "fournisseur_nom": "Mouser",
            "type": "Devis",
            "statut": "Devis recu",
            "date_reception_devis": "2026-10-05",
        },
    ).json()["numero"]
    for composant in composants:
        ligne = {"composant_id": composant, "qte_commandee": 2, "pu_ht_devis": 1.25}
        client.post(f"/api/commandes/{numero}/lignes", json=ligne)
    client.post(
        f"/api/commandes/{numero}/frais", json={"type_frais": "Livraison", "montant_ht": 9.9}
    )
    if transmise:
        client.patch(f"/api/commandes/{numero}", json={"date_transmission_suivi": "2026-10-06"})
    return numero


def _appliquer(
    client: TestClient,
    contenu: bytes,
    retenues: dict | None = None,
    associations: list[str] | None = None,
) -> bytes:
    donnees = {"retenues": json.dumps(retenues or {})}
    if associations is not None:
        donnees["associations"] = json.dumps(associations)
    reponse = client.post(
        "/api/suivi-achats/appliquer", files={"fichier": ("SUIVI.xlsx", contenu)}, data=donnees
    )
    assert reponse.status_code == 200, reponse.text
    assert "SUIVI.xlsx" in reponse.headers["content-disposition"]
    return reponse.content


def _analyse(client: TestClient, contenu: bytes) -> dict:
    reponse = client.post("/api/suivi-achats/analyse", files={"fichier": ("SUIVI.xlsx", contenu)})
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def test_aller_ajoute_les_devis_transmis(client_essai: TestClient) -> None:
    premier = _commande(client_essai, ["ESSAI-ALI-001", "ESSAI-ALI-002", "ESSAI-ALI-003"])
    second = _commande(client_essai, ["ESSAI-AGI-001", "ESSAI-AGI-002"])
    _commande(client_essai, ["ESSAI-OBS-001"], transmise=False)
    client_essai.patch("/api/composants/ESSAI-ALI-002", json={"designation": "=HYPERLINK(1)"})
    assert _analyse(client_essai, _modele())["nb_lignes_a_ajouter"] == 7
    feuille = load_workbook(io.BytesIO(_appliquer(client_essai, _modele()))).active
    references = [feuille.cell(r, 12).value for r in range(10, 17)]
    assert references[0] == f"{premier} · ESSAI-ALI-001"
    assert f"{premier} · frais 1" in references and f"{second} · ESSAI-AGI-002" in references
    assert feuille.column_dimensions["L"].hidden
    # Une ligne par composant : fournisseur, référence, quantité, montant, date du devis.
    assert feuille["A10"].value == "Mouser" and feuille["D10"].value == 2
    assert feuille["E10"].value == 2.5 and feuille["F10"].value == datetime(2026, 10, 5)
    # Une désignation qui commence par « = » reste du texte.
    assert feuille["C11"].data_type == "s"
    # Plus de lignes que le modèle n'en prévoit : le TOTAL descend, ses sommes s'étendent.
    assert feuille["D17"].value == "TOTAL"
    assert feuille["E17"].value == "=SUM(E10:E16)" and feuille["J17"].value == "=SUM(J10:J16)"
    assert feuille["A20"].value.startswith("Colonnes à droite")


def test_aller_n_ajoute_rien_deux_fois(client_essai: TestClient) -> None:
    _commande(client_essai, ["ESSAI-ALI-001"])
    une_fois = _appliquer(client_essai, _modele())
    deux_fois = load_workbook(io.BytesIO(_appliquer(client_essai, une_fois))).active
    remplies = [r for r in range(10, 20) if deux_fois.cell(r, 12).value]
    assert len(remplies) == 2
    assert _analyse(client_essai, une_fois)["nb_lignes_a_ajouter"] == 0


def test_retour_propose_puis_applique_ce_qui_est_retenu(client_essai: TestClient) -> None:
    numero = _commande(client_essai, ["ESSAI-ALI-001", "ESSAI-ALI-002"])
    classeur = load_workbook(io.BytesIO(_appliquer(client_essai, _modele())))
    feuille = classeur.active
    # L'établissement colore la commande en « commandé », date la commande (bleu :
    # confirmée), annonce une livraison et saisit le montant réel sur deux lignes.
    for rang in (10, 11, 12):
        for colonne in range(1, 12):
            feuille.cell(rang, colonne).fill = COMMANDE
    feuille["G10"].value, feuille["G10"].font = datetime(2026, 10, 8), BLEU
    feuille["H10"].value = datetime(2026, 10, 20)
    feuille["J10"].value, feuille["J11"].value = 30, 4.5
    feuille["K10"].value = "Bon de commande signé"
    feuille["A13"].value, feuille["C13"].value = "Autre", "Ligne saisie à la main"
    flux = io.BytesIO()
    classeur.save(flux)
    analyse = _analyse(client_essai, flux.getvalue())
    assert analyse["lignes_inconnues"] == ["ligne 13 : Autre · Ligne saisie à la main"]
    changements = {c["champ"]: c["apres"] for c in analyse["propositions"][0]["changements"]}
    assert changements == {
        "type": "Commande",
        "statut": "Commande",
        "date_commande": "2026-10-08",
        "date_commande_confirmee": 1,
        "livraison_annoncee": "2026-10-20",
        "montant_commande_ht": 34.5,
        "commentaire_suivi": "Bon de commande signé",
    }
    retenus = {numero: ["type", "statut", "date_commande", "montant_commande_ht"]}
    _appliquer(client_essai, flux.getvalue(), retenus)
    commande = client_essai.get(f"/api/commandes/{numero}").json()
    assert commande["statut"] == "Commande" and commande["date_commande"] == "2026-10-08"
    assert commande["livraison_annoncee"] is None  # non retenu
    # Le montant réel remplace l'estimation (2 × 2 × 1,25 + 9,90 = 14,90) dans l'engagé.
    assert client_essai.get("/api/pilotage").json()["montant_engage_ht"] == 34.5


def test_livre_par_la_couleur_n_entraine_pas_de_reception(client_essai: TestClient) -> None:
    numero = _commande(client_essai, ["ESSAI-ALI-001"])
    client_essai.patch(f"/api/commandes/{numero}", json={"type": "Commande", "statut": "Commande"})
    classeur = load_workbook(io.BytesIO(_appliquer(client_essai, _modele())))
    for rang in (10, 11):
        for colonne in range(1, 12):
            classeur.active.cell(rang, colonne).fill = LIVRE
    flux = io.BytesIO()
    classeur.save(flux)
    changements = _analyse(client_essai, flux.getvalue())["propositions"][0]["changements"]
    assert [c["champ"] for c in changements] == ["date_reception_annoncee"]
    _appliquer(client_essai, flux.getvalue(), {numero: ["date_reception_annoncee"]})
    commande = client_essai.get(f"/api/commandes/{numero}").json()
    assert commande["statut"] == "Commande" and commande["date_reception_annoncee"]


def test_classeur_sans_le_modele_refuse(client_essai: TestClient) -> None:
    classeur = Workbook()
    classeur.active["A1"] = "Autre chose"
    flux = io.BytesIO()
    classeur.save(flux)
    reponse = client_essai.post(
        "/api/suivi-achats/analyse", files={"fichier": ("x.xlsx", flux.getvalue())}
    )
    assert reponse.status_code == 400
    assert "modèle de suivi" in reponse.json()["erreur"]


def _sans_date_de_devis(client: TestClient) -> str:
    numero = client.post(
        "/api/commandes",
        json={"fournisseur_nom": "Mouser", "type": "Devis", "statut": "Devis demande"},
    ).json()["numero"]
    ligne = {"composant_id": "ESSAI-ALI-001", "qte_commandee": 2, "pu_ht_devis": 1.25}
    client.post(f"/api/commandes/{numero}/lignes", json=ligne)
    client.patch(f"/api/commandes/{numero}", json={"date_transmission_suivi": "2026-10-06"})
    return numero


def test_case_vide_completee_au_passage_suivant(client_essai: TestClient) -> None:
    """Transmise sans date de devis ; la date saisie ensuite remplit la case vide de la ligne
    déjà écrite, et une case remplie à la main n'est jamais réécrite."""
    numero = _sans_date_de_devis(client_essai)
    premier = load_workbook(io.BytesIO(_appliquer(client_essai, _modele())))
    feuille = premier.active
    assert feuille["F10"].value is None
    feuille["A10"] = "Mouser France"  # retouché à la main dans le fichier
    flux = io.BytesIO()
    premier.save(flux)
    client_essai.patch(f"/api/commandes/{numero}", json={"date_reception_devis": "2026-10-09"})
    assert _analyse(client_essai, flux.getvalue())["nb_lignes_a_completer"] == 1
    second = load_workbook(io.BytesIO(_appliquer(client_essai, flux.getvalue()))).active
    assert second["F10"].value == datetime(2026, 10, 9)
    assert second["A10"].value == "Mouser France"
    assert [r for r in range(10, 20) if second.cell(r, 12).value] == [10]


def test_date_du_devis_au_statut_devis_recu(client_essai: TestClient) -> None:
    numero = _sans_date_de_devis(client_essai)
    client_essai.patch(f"/api/commandes/{numero}", json={"statut": "Devis recu"})
    assert client_essai.get(f"/api/commandes/{numero}").json()["date_reception_devis"] == (
        date.today().isoformat()
    )
    # Une date déjà saisie n'est pas remplacée.
    autre = _commande(client_essai, ["ESSAI-ALI-002"])
    client_essai.patch(f"/api/commandes/{autre}", json={"statut": "Devis demande"})
    client_essai.patch(f"/api/commandes/{autre}", json={"statut": "Devis recu"})
    assert (
        client_essai.get(f"/api/commandes/{autre}").json()["date_reception_devis"] == "2026-10-05"
    )


def test_date_du_devis_au_depot_d_un_devis(client_essai: TestClient) -> None:
    numero = _sans_date_de_devis(client_essai)

    def deposer(type_document: str) -> None:
        reponse = client_essai.post(
            f"/api/commandes/{numero}/documents",
            files=[("fichiers", (f"{type_document}.pdf", f"%PDF {type_document}".encode()))],
            data={"type_document": type_document},
        )
        assert reponse.status_code == 201, reponse.text

    deposer("Facture")
    assert client_essai.get(f"/api/commandes/{numero}").json()["date_reception_devis"] is None
    deposer("Devis")
    assert client_essai.get(f"/api/commandes/{numero}").json()["date_reception_devis"] == (
        date.today().isoformat()
    )


# --- Cellules fusionnées et lignes saisies à la main (ACH-05) ------------------------------------


def _enregistrer(classeur: Workbook) -> bytes:
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _fusionner_commande(feuille, premier: int, dernier: int) -> None:
    """Ce que fait l'établissement : une case par commande pour le fournisseur, les dates et
    le montant de la commande, fusionnée sur toutes ses lignes."""
    for colonne in ("A", "F", "G", "H", "J"):
        feuille.merge_cells(f"{colonne}{premier}:{colonne}{dernier}")


def test_cellules_fusionnees_par_l_etablissement(client_essai: TestClient) -> None:
    """Le 28/10 : erreur 500 en complétant une case vide qui était une cellule fusionnée."""
    numero = _commande(client_essai, ["ESSAI-ALI-001", "ESSAI-ALI-003"])
    classeur = load_workbook(io.BytesIO(_appliquer(client_essai, _modele())))
    feuille = classeur.active
    feuille["F10"].value = None  # date du devis effacée : une case à compléter…
    _fusionner_commande(feuille, 10, 12)  # … qui n'existe plus que dans le bloc fusionné
    feuille["F10"].value = datetime(2026, 10, 5)
    feuille["J10"].value = 99.0
    contenu = _enregistrer(classeur)
    changements = _analyse(client_essai, contenu)["propositions"][0]["changements"]
    # Le montant fusionné compte une fois, pas une fois par ligne du bloc.
    assert {c["champ"]: c["apres"] for c in changements}["montant_commande_ht"] == 99.0
    _commande(client_essai, ["ESSAI-AGI-003"])  # une commande de plus : des lignes à ajouter
    rendu = load_workbook(io.BytesIO(_appliquer(client_essai, contenu, {numero: []}))).active
    assert {str(r) for r in rendu.merged_cells.ranges} >= {"A10:A12", "F10:F12", "J10:J12"}
    assert rendu["J10"].value == 99.0
    assert [rendu.cell(r, 12).value for r in (13, 14)] == [
        "CMD-002 · ESSAI-AGI-003",
        "CMD-002 · frais 2",
    ]


def test_lignes_ajoutees_decalent_les_fusions_sous_le_total(client_essai: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(_modele()))
    feuille = classeur.active
    feuille.merge_cells("G16:H16")  # une note fusionnée sous le TOTAL, comme la légende
    feuille["G16"].value = "En bleu, date confirmée"
    _commande(client_essai, ["ESSAI-ALI-001", "ESSAI-ALI-002", "ESSAI-ALI-003", "ESSAI-AGI-003"])
    rendu = load_workbook(io.BytesIO(_appliquer(client_essai, _enregistrer(classeur)))).active
    # 5 lignes pour 4 places : une ligne insérée, la note descend d'autant avec sa fusion.
    assert rendu["G17"].value == "En bleu, date confirmée"
    assert "G17:H17" in {str(r) for r in rendu.merged_cells.ranges}


def _saisie_a_la_main(feuille, lignes: list[tuple]) -> None:
    """Lignes tapées par l'établissement, sans référence Nomentrace, à partir de la ligne 10."""
    for decalage, (fournisseur, reference, designation, qte, montant) in enumerate(lignes):
        rang = 10 + decalage
        for colonne, valeur in enumerate((fournisseur, reference, designation, qte, montant), 1):
            if valeur is not None:
                feuille.cell(rang, colonne, valeur)


def _commande_ecrite_a_la_main(client: TestClient, qte_lue: int = 2) -> tuple[str, bytes]:
    numero = _commande(client, ["ESSAI-ALI-003", "ESSAI-AGI-003", "ESSAI-ALI-002"])
    classeur = load_workbook(io.BytesIO(_modele()))
    feuille = classeur.active
    _saisie_a_la_main(
        feuille,
        [
            (" MOUSER", "ddr 60l-5", "Convertisseur 5 V", 2, 2.5),  # même référence
            (None, "-", "Codeur magnetique AS5047D ATSM", 2, 2.5),  # désignation proche
            (None, None, "Carte de puissance custom", qte_lue, 2.5),
            (None, None, "Frais de port", None, 9.9),  # le frais, par son montant
        ],
    )
    _fusionner_commande(feuille, 10, 13)
    feuille["G10"].value = datetime(2026, 11, 2)
    return numero, _enregistrer(classeur)


def test_lignes_saisies_a_la_main_associees(client_essai: TestClient) -> None:
    numero, contenu = _commande_ecrite_a_la_main(client_essai)
    analyse = _analyse(client_essai, contenu)
    assert analyse["lignes_inconnues"] == []
    assert analyse["nb_lignes_a_ajouter"] == 0
    association = analyse["associations"][0]
    assert (association["numero"], association["complete"]) == (numero, True)
    assert [ligne["ligne"] for ligne in association["lignes"]] == [10, 11, 12, 13]
    assert all(not ligne["ecarts"] for ligne in association["lignes"])
    # La commande associée reçoit les propositions de ses lignes, date de commande comprise.
    changements = {c["champ"] for c in analyse["propositions"][0]["changements"]}
    assert "date_commande" in changements
    rendu = load_workbook(
        io.BytesIO(_appliquer(client_essai, contenu, associations=[numero]))
    ).active
    assert [rendu.cell(r, 12).value for r in range(10, 14)] == [
        f"{numero} · ESSAI-ALI-003",
        f"{numero} · ESSAI-AGI-003",
        f"{numero} · ESSAI-ALI-002",
        f"{numero} · frais 1",
    ]
    assert rendu.cell(14, 12).value is None  # rien d'ajouté une seconde fois
    assert (
        _analyse(
            client_essai,
            _enregistrer(
                load_workbook(io.BytesIO(_appliquer(client_essai, contenu, associations=[numero])))
            ),
        )["associations"]
        == []
    )


def test_association_non_retenue_ajoute_les_lignes(client_essai: TestClient) -> None:
    numero, contenu = _commande_ecrite_a_la_main(client_essai)
    rendu = load_workbook(io.BytesIO(_appliquer(client_essai, contenu))).active
    assert rendu.cell(10, 12).value is None
    assert rendu.cell(14, 12).value == f"{numero} · ESSAI-ALI-003"


def test_association_signale_l_ecart_de_quantite(client_essai: TestClient) -> None:
    _, contenu = _commande_ecrite_a_la_main(client_essai, qte_lue=5)
    association = _analyse(client_essai, contenu)["associations"][0]
    ecarts = {ligne["ligne"]: ligne["ecarts"] for ligne in association["lignes"]}
    assert ecarts[12] == ["quantité 5 dans le fichier, 2 dans Nomentrace"]


def test_ligne_inconnue_decrite(client_essai: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(_modele()))
    _saisie_a_la_main(classeur.active, [("Autre", None, "Ligne saisie à la main", 1, 3.0)])
    analyse = _analyse(client_essai, _enregistrer(classeur))
    assert analyse["lignes_inconnues"] == ["ligne 10 : Autre · Ligne saisie à la main"]
