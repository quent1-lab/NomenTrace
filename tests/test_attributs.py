"""Attributs paramétrables : typage, saisie, filtre et tri, export, import, répartition."""

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_import import _appliquer, _decisions_par_defaut, _deposer, _lignes, _modifier

C1, C2, C3 = "ESSAI-ALI-001", "ESSAI-ALI-003", "ESSAI-ALI-006"


def _creer_attribut(client: TestClient, libelle: str, type_: str, unite: str | None = None) -> dict:
    reponse = client.post(
        "/api/attributs", json={"libelle": libelle, "type": type_, "unite": unite}
    )
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


def _saisir(client: TestClient, composant: str, **valeurs: Any):
    return client.put(f"/api/composants/{composant}/attributs", json={"valeurs": valeurs})


@pytest.fixture
def client_attributs(client_essai: TestClient) -> TestClient:
    """Tension (nombre, V), Matériau (liste), Étanche (oui/non) ; trois composants renseignés."""
    _creer_attribut(client_essai, "Tension", "nombre", "V")
    _creer_attribut(client_essai, "Matériau", "liste")
    for libelle in ("Acier", "Inox", "Laiton"):
        client_essai.post("/api/attributs/materiau/valeurs", json={"libelle": libelle})
    _creer_attribut(client_essai, "Étanche", "booleen")
    assert _saisir(client_essai, C1, tension=48, materiau="Acier").status_code == 200
    assert _saisir(client_essai, C2, tension="12", etanche="oui").status_code == 200
    assert _saisir(client_essai, C3, tension="3,3", materiau="Inox").status_code == 200
    return client_essai


# --- Création et typage --------------------------------------------------------------------------


def test_creation_et_typage(client_essai: TestClient) -> None:
    tension = _creer_attribut(client_essai, "Tension nominale", "nombre", " V ")
    assert tension["code"] == "tension_nominale"
    assert tension["unite"] == "V"
    assert tension["type"] == "nombre"
    materiau = _creer_attribut(client_essai, "Matériau", "liste")
    assert materiau["code"] == "materiau"
    assert (
        client_essai.post(
            "/api/attributs", json={"libelle": "materiau", "type": "texte"}
        ).status_code
        == 409
    )
    assert (
        client_essai.post("/api/attributs", json={"libelle": "Poids", "type": "date"}).status_code
        == 422
    )
    # Le type est figé : il ne fait pas partie des champs modifiables.
    assert client_essai.patch("/api/attributs/materiau", json={"type": "texte"}).status_code == 422
    renomme = client_essai.patch("/api/attributs/materiau", json={"libelle": "Matière"}).json()
    assert (renomme["code"], renomme["libelle"]) == ("materiau", "Matière")


def test_valeurs_de_liste(client_essai: TestClient) -> None:
    _creer_attribut(client_essai, "Matériau", "liste")
    attribut = client_essai.post(
        "/api/attributs/materiau/valeurs", json={"libelle": "Aluminium"}
    ).json()
    assert [(v["code"], v["libelle"]) for v in attribut["valeurs"]] == [("Aluminium", "Aluminium")]
    renomme = client_essai.patch(
        "/api/attributs/materiau/valeurs/Aluminium", json={"libelle": "Alu 6061"}
    )
    assert renomme.json()["valeurs"][0]["code"] == "Aluminium"
    _creer_attribut(client_essai, "Tension", "nombre")
    refus = client_essai.post("/api/attributs/tension/valeurs", json={"libelle": "12"})
    assert refus.status_code == 400


def test_saisie_valeurs_et_journal(client_attributs: TestClient) -> None:
    fiche = client_attributs.get(f"/api/composants/{C3}").json()
    assert fiche["attributs"] == {"tension": 3.3, "materiau": "Inox"}
    assert any(
        j["table_cible"] == "composant_attribut" and j["champ"] == "tension"
        for j in fiche["journal"]
    )
    assert client_attributs.get(f"/api/composants/{C2}").json()["attributs"]["etanche"] == "1"
    # null efface la valeur
    assert _saisir(client_attributs, C3, materiau=None).json() == {"tension": 3.3}


@pytest.mark.parametrize(
    ("valeurs", "message"),
    [
        ({"tension": "douze"}, "n'est pas un nombre"),
        ({"materiau": "Bois"}, "n'est pas une valeur de la liste"),
        ({"etanche": "peut-être"}, "ni oui ni non"),
        ({"inconnu": "x"}, "inconnu"),
    ],
)
def test_valeur_refusee(client_attributs: TestClient, valeurs: dict, message: str) -> None:
    reponse = _saisir(client_attributs, C1, **valeurs)
    assert reponse.status_code in (400, 404)
    assert message in reponse.json()["erreur"]
    assert client_attributs.get(f"/api/composants/{C1}").json()["attributs"]["tension"] == 48


def test_valeur_ou_attribut_desactive(client_attributs: TestClient) -> None:
    client_attributs.patch("/api/attributs/materiau/valeurs/Inox", json={"actif": 0})
    assert _saisir(client_attributs, C1, materiau="Inox").status_code == 400
    # Déjà portée par le composant : elle reste acceptée telle quelle.
    assert _saisir(client_attributs, C3, materiau="Inox").status_code == 200
    client_attributs.patch("/api/attributs/tension", json={"actif": 0})
    assert _saisir(client_attributs, C1, tension=24).status_code == 400


# --- Filtre et tri -------------------------------------------------------------------------------


def _ids(client: TestClient, **params: Any) -> list[str]:
    reponse = client.get("/api/composants", params=params)
    assert reponse.status_code == 200, reponse.text
    return [c["id"] for c in reponse.json()]


def test_filtres_sur_attribut(client_attributs: TestClient) -> None:
    assert _ids(client_attributs, attr="tension:egal:12") == [C2]
    assert _ids(client_attributs, attr="tension:egal:3.3 V") == [C3]
    assert _ids(client_attributs, attr="tension:entre:5:50") == [C1, C2]
    assert _ids(client_attributs, attr="tension:entre::5") == [C3]
    assert _ids(client_attributs, attr="materiau:egal:Inox") == [C3]
    assert _ids(client_attributs, attr="etanche:egal:oui") == [C2]
    assert _ids(client_attributs, attr="materiau:renseigne") == [C1, C3]
    assert len(_ids(client_attributs, attr="tension:vide")) == 57
    # Deux filtres se cumulent.
    assert _ids(client_attributs, attr=["tension:renseigne", "materiau:vide"]) == [C2]
    assert _ids(client_attributs, attr="tension:renseigne", bloc="ALI", q="003") == [C2]


def test_tri_sur_attribut(client_attributs: TestClient) -> None:
    montant = _ids(client_attributs, tri="attr:tension")
    assert montant[:3] == [C3, C2, C1]
    assert len(montant) == 60  # sans valeur : en fin de liste
    assert _ids(client_attributs, tri="attr:tension", ordre="desc")[:3] == [C1, C2, C3]
    # Une liste se trie dans l'ordre de ses valeurs (Acier, Inox, Laiton).
    assert _ids(client_attributs, tri="attr:materiau")[:2] == [C1, C3]
    ligne = client_attributs.get("/api/composants", params={"q": C3}).json()[0]
    assert ligne["attributs"] == {"tension": 3.3, "materiau": "Inox"}


@pytest.mark.parametrize(
    "params",
    [
        {"tri": "attr:inconnu"},
        {"tri": "attr:tension; DROP TABLE composant"},
        {"attr": "inconnu:vide"},
        {"attr": "tension:plus_grand:3"},
        {"attr": "materiau:entre:1:2"},
        {"attr": "materiau:egal:Bois"},
        {"attr": "tension:entre:a:b"},
    ],
)
def test_filtre_ou_tri_refuse(client_attributs: TestClient, params: dict) -> None:
    reponse = client_attributs.get("/api/composants", params=params)
    assert reponse.status_code == 400
    assert "erreur" in reponse.json()
    assert len(_ids(client_attributs)) == 60


# --- Exports -------------------------------------------------------------------------------------


def _feuille(contenu: bytes) -> list[tuple]:
    classeur = load_workbook(io.BytesIO(contenu))
    return [tuple(ligne) for ligne in classeur["Composants"].iter_rows(values_only=True)]


def test_export_filtre_avec_colonnes_d_attributs(client_attributs: TestClient) -> None:
    params = {"colonnes": "id,attr:tension,attr:materiau,attr:etanche", "attr": "tension:renseigne"}
    lignes = _feuille(client_attributs.get("/api/composants/export", params=params).content)
    assert lignes[0] == ("ID", "Tension (V)", "Matériau", "Étanche")
    assert lignes[1:4] == [(C1, 48, "Acier", None), (C2, 12, None, "Oui"), (C3, 3.3, "Inox", None)]
    toutes = _feuille(client_attributs.get("/api/composants/export").content)
    assert toutes[0][-3:] == ("Tension (V)", "Matériau", "Étanche")
    refus = client_attributs.get("/api/composants/export", params={"colonnes": "id,attr:inconnu"})
    assert refus.status_code == 400


def test_excel_global_avec_attributs(client_attributs: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(client_attributs.get("/api/export/classeur").content))
    feuille = classeur["composant"]
    entetes = [c.value for c in feuille[1]]
    assert entetes[-3:] == ["Tension (V)", "Matériau", "Étanche"]
    ligne = next(r for r in feuille.iter_rows(min_row=2, values_only=True) if r[0] == C3)
    assert ligne[-3:] == (3.3, "Inox", None)
    assert "composant_attribut" in classeur.sheetnames


# --- Import --------------------------------------------------------------------------------------


def test_modele_porte_les_attributs(client_attributs: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(client_attributs.get("/api/blocs/ALI/modele").content))
    feuille = classeur["Composants"]
    entetes = [c.value for c in feuille[1]]
    colonne = entetes.index("Matériau")
    assert "Tension (V)" in entetes and "Étanche" in entetes
    ligne = next(r for r in feuille.iter_rows(min_row=2, values_only=True) if r[0] == C1)
    assert (ligne[entetes.index("Tension (V)")], ligne[colonne]) == (48, "Acier")
    lettre = feuille.cell(1, colonne + 1).column_letter
    validation = next(
        v for v in feuille.data_validations.dataValidation if f"{lettre}2" in str(v.sqref)
    )
    assert validation.errorStyle in (None, "stop")
    # Redéposé tel quel, le modèle ne propose aucune modification.
    depot = _deposer(
        client_attributs, [("ALI.xlsx", client_attributs.get("/api/blocs/ALI/modele").content)]
    )
    assert depot["resume"] == {"IDENTIQUE": 20}


def test_import_compare_et_applique_les_attributs(client_attributs: TestClient) -> None:
    contenu = client_attributs.get("/api/blocs/ALI/modele").content
    contenu = _modifier(contenu, C1, "Tension (V)", 24)
    contenu = _modifier(contenu, C1, "Matériau", "laiton")
    contenu = _modifier(contenu, C3, "Tension (V)", None)
    classeur = load_workbook(io.BytesIO(contenu))
    feuille = classeur["Composants"]
    entetes = [c.value for c in feuille[1]]
    nouvelle = feuille.max_row + 1
    for entete, valeur in {
        "Fonction": "Mesure",
        "Désignation": "Sonde de température",
        "Mode appro": "Achat",
        "Qté besoin": 2,
        "Tension (V)": "5 V",
        "Étanche": "Oui",
    }.items():
        feuille.cell(nouvelle, entetes.index(entete) + 1, valeur)
    sortie = io.BytesIO()
    classeur.save(sortie)

    depot = _deposer(client_attributs, [("ALI.xlsx", sortie.getvalue())])

    modifies = {
        ligne["composant_id"]: ligne["donnees"]["differences"]
        for ligne in _lignes(depot, "MODIFIE")
    }
    assert modifies[C1] == {
        "attr:tension": {"actuel": 48, "propose": 24},
        "attr:materiau": {"actuel": "Acier", "propose": "Laiton"},
    }
    assert modifies[C3] == {"attr:tension": {"actuel": 3.3, "propose": None}}
    nouveau = _lignes(depot, "NOUVEAU")[0]["donnees"]["valeurs"]
    assert (nouveau["attr:tension"], nouveau["attr:etanche"]) == (5, "1")

    resultat = _appliquer(client_attributs, depot, _decisions_par_defaut(depot))

    assert resultat["refus"] == []
    assert client_attributs.get(f"/api/composants/{C1}").json()["attributs"] == {
        "tension": 24,
        "materiau": "Laiton",
    }
    assert client_attributs.get(f"/api/composants/{C3}").json()["attributs"] == {"materiau": "Inox"}
    cree = client_attributs.get(f"/api/composants/{resultat['crees'][0]}").json()
    assert cree["attributs"] == {"tension": 5, "etanche": "1"}
    assert any(
        j["origine"] == "import"
        for j in cree["journal"]
        if j["table_cible"] == "composant_attribut"
    )


def test_import_valeur_d_attribut_invalide(client_attributs: TestClient) -> None:
    contenu = _modifier(
        client_attributs.get("/api/blocs/ALI/modele").content, C1, "Matériau", "Bois"
    )
    depot = _deposer(client_attributs, [("ALI.xlsx", contenu)])
    invalide = _lignes(depot, "INVALIDE")[0]
    assert invalide["composant_id"] == C1
    assert "n'est pas une valeur de la liste" in invalide["donnees"]["raisons"][0]


# --- Répartition, suppression, reclassement ------------------------------------------------------


def test_repartition_des_valeurs(client_attributs: TestClient) -> None:
    _saisir(client_attributs, "ESSAI-TR-001", tension=12)
    reponse = client_attributs.get("/api/attributs/tension/repartition").json()
    assert [(v["valeur"], v["nb_composants"]) for v in reponse["valeurs"]] == [
        (3.3, 1),
        (12, 2),
        (48, 1),
    ]
    assert reponse["non_renseigne"]["nb_composants"] == 56
    douze = next(v for v in reponse["valeurs"] if v["valeur"] == 12)
    assert douze["cout_ht"] == 38.52  # ALI-003 (34 €) et TR-001 (4 × 1,13 €)
    filtre = client_attributs.get(
        "/api/attributs/tension/repartition", params={"bloc": "TR"}
    ).json()
    assert [(v["valeur"], v["nb_composants"]) for v in filtre["valeurs"]] == [(12, 1)]
    materiau = client_attributs.get("/api/attributs/materiau/repartition").json()
    assert [v["libelle"] for v in materiau["valeurs"]] == ["Acier", "Inox"]
    assert client_attributs.get("/api/attributs/inconnu/repartition").status_code == 404


def test_suppression_et_reclassement_gardent_la_coherence(client_attributs: TestClient) -> None:
    nouveau = client_attributs.post(
        f"/api/composants/{C1}/reclasser", json={"bloc_code": "TR"}
    ).json()
    assert client_attributs.get(f"/api/composants/{nouveau['id']}").json()["attributs"] == {
        "tension": 48,
        "materiau": "Acier",
    }
    corps = {"cles": [C3], "action": "supprimer"}
    assert client_attributs.post("/api/nettoyage/composants", json=corps).json()["supprimes"] == [
        C3
    ]
