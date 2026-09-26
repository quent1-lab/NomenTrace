"""Téléchargements : export filtré des composants, Excel global et archive complète."""

import io
import re
import sqlite3
import tempfile
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

PDF = b"%PDF-1.4\n% fiche technique de test\n"
MOTIF_NOM = r'attachment; filename="nomentrace_composants_\d{8}_\d{4}\.xlsx"'


def _feuille(contenu: bytes) -> list[tuple]:
    classeur = load_workbook(io.BytesIO(contenu))
    return [tuple(ligne) for ligne in classeur["Composants"].iter_rows(values_only=True)]


def test_export_filtre_colonnes_lignes_totaux(client_essai: TestClient) -> None:
    filtres = {"bloc": "ALI", "tri": "total_ht", "ordre": "desc"}
    colonnes = "id,avancement,pu_releve,designation,statut_appro,total_ht"
    reponse = client_essai.get("/api/composants/export", params={**filtres, "colonnes": colonnes})
    assert reponse.status_code == 200
    assert re.fullmatch(MOTIF_NOM, reponse.headers["content-disposition"])
    lignes = _feuille(reponse.content)
    assert lignes[0] == ("ID", "Avancement", "PU relevé", "Désignation", "Statut appro", "Total HT")
    ecran = client_essai.get("/api/composants", params=filtres).json()
    assert [ligne[0] for ligne in lignes[1:-1]] == [c["id"] for c in ecran]
    assert len(ecran) == 20
    assert lignes[-1] == ("20 composant(s) affiché(s)", None, None, None, "Total HT", 447.27)


def test_export_filtre_valeurs_en_libelle(client_essai: TestClient) -> None:
    colonnes = "id,pu_releve,pu_ht,statut_appro,avancement"
    reponse = client_essai.get(
        "/api/composants/export", params={"bloc": "ALI", "colonnes": colonnes}
    )
    classeur = load_workbook(io.BytesIO(reponse.content))
    feuille = classeur["Composants"]
    par_id = {ligne[0].value: ligne for ligne in feuille.iter_rows(min_row=2)}
    ttc = par_id["ESSAI-ALI-014"]
    assert ttc[1].value == 34.0
    assert "TTC" in ttc[1].number_format
    assert ttc[2].value == 28.33
    assert ttc[3].value == "Non lancé"
    assert ttc[4].value == "À commander"
    assert par_id["ESSAI-ALI-007"][1].value == "à chiffrer"
    assert par_id["ESSAI-ALI-002"][1].value is None
    assert par_id["ESSAI-ALI-002"][4].value == "Hors achat"


def test_export_libelle_des_listes_parametrables(client_essai: TestClient) -> None:
    reponse = client_essai.get(
        "/api/composants/export",
        params={"mode_appro": "Fourni partenaire", "colonnes": "id,mode_appro"},
    )
    lignes = _feuille(reponse.content)
    assert len(lignes) > 2
    assert {ligne[1] for ligne in lignes[1:-1]} == {"Fourni par le partenaire"}


def test_export_toutes_colonnes_par_defaut(client_essai: TestClient) -> None:
    lignes = _feuille(client_essai.get("/api/composants/export").content)
    assert lignes[0][:4] == ("ID", "Bloc", "Fonction", "Désignation")
    assert len(lignes[0]) == 20
    assert len(lignes) == 60 + 2
    assert lignes[-1][0] == "60 composant(s) affiché(s)"
    assert lignes[-1][lignes[0].index("Total HT")] == 3184.98


def test_export_colonne_hors_liste_blanche(client_essai: TestClient) -> None:
    reponse = client_essai.get("/api/composants/export", params={"colonnes": "id,note_technique"})
    assert reponse.status_code == 400
    assert "note_technique" in reponse.json()["erreur"]


def test_export_filtre_rien_dans_echange(client_essai: TestClient, dossier_echange: Path) -> None:
    client_essai.get("/api/composants/export", params={"bloc": "ALI"})
    client_essai.get("/api/export/classeur")
    assert list(dossier_echange.rglob("nomentrace_*.xlsx")) == []


def test_telechargement_excel_global(client_essai: TestClient) -> None:
    reponse = client_essai.get("/api/export/classeur")
    assert reponse.status_code == 200
    assert re.search(r"nomentrace_export_\d{8}_\d{4}\.xlsx", reponse.headers["content-disposition"])
    classeur = load_workbook(io.BytesIO(reponse.content), read_only=True)
    assert classeur.sheetnames[:4] == ["Pilotage", "Blocs", "Ensembles", "Affectations"]
    pilotage = classeur["Pilotage"]
    valeurs = {ligne[0]: ligne[1] for ligne in pilotage.iter_rows(min_row=4, values_only=True)}
    assert valeurs["cout_ht"] == 3184.98
    classeur.close()


def test_archive_complete_relisible(
    client_essai: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    depot = client_essai.post(
        "/api/composants/ESSAI-ALI-001/documents",
        files=[("fichiers", ("Fiche alim.pdf", PDF))],
        data={"type_document": "Fiche technique"},
    )
    assert depot.status_code == 201
    temporaire = tmp_path / "temporaire"
    temporaire.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(temporaire))

    reponse = client_essai.get("/api/sauvegardes/archive")
    assert reponse.status_code == 200
    assert re.search(r"nomentrace_archive_\d{8}_\d{4}\.zip", reponse.headers["content-disposition"])
    assert list(temporaire.iterdir()) == []  # fichier temporaire supprimé après envoi

    extrait = tmp_path / "extrait"
    with zipfile.ZipFile(io.BytesIO(reponse.content)) as archive:
        # La base des comptes ne sort jamais par l'interface, seulement par le serveur.
        assert "comptes.db" not in archive.namelist()
        archive.extractall(extrait)
    documents = [p for p in (extrait / "documents").rglob("*") if p.is_file()]
    assert [p.read_bytes() for p in documents] == [PDF]
    conn = sqlite3.connect(extrait / "nomentrace.db")
    try:
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert conn.execute("SELECT COUNT(*) FROM composant").fetchone()[0] == 60
        chemin = conn.execute("SELECT chemin FROM document").fetchone()[0]
    finally:
        conn.close()
    assert (extrait / "documents" / chemin).read_bytes() == PDF
