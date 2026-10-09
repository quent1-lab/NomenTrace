"""Document de demande de devis envoyé au fournisseur, en PDF et en Excel."""

from io import BytesIO
from urllib.parse import unquote

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import db
from backend.services import demande_document
from backend.services.droits import Utilisateur


def _demande(client: TestClient) -> str:
    client.patch(
        "/api/parametres",
        json={
            "etablissement_nom": "École d'essai",
            "adresse_livraison": "1 rue des Essais\n75000 Paris",
            "contact_achats_nom": "Contact Achats",
            "contact_achats_mail": "achats@ecole.test",
        },
    )
    numero = client.post("/api/commandes", json={"fournisseur_nom": "Mouser"}).json()["numero"]
    for composant in ("ESSAI-ALI-001", "ESSAI-AGI-001", "ESSAI-OBS-001"):
        ligne = {"composant_id": composant, "qte_commandee": 2}
        client.post(f"/api/commandes/{numero}/lignes", json=ligne)
    return numero


def test_excel_de_la_demande(client_essai: TestClient) -> None:
    numero = _demande(client_essai)
    reponse = client_essai.get(f"/api/commandes/{numero}/demande-devis.xlsx?avant=2026-11-02")
    assert reponse.status_code == 200
    disposition = unquote(reponse.headers["content-disposition"])
    assert f"Demande de devis {numero} Mouser.xlsx" in disposition
    feuille = load_workbook(BytesIO(reponse.content)).active
    valeurs = [c.value for ligne in feuille.iter_rows() for c in ligne if c.value is not None]
    assert "École d'essai" in valeurs and "02/11/2026" in valeurs
    # Le demandeur en mode local, faute de compte : le contact achats.
    assert "Contact Achats · achats@ecole.test" in valeurs
    assert "Alimentation AC/DC 48 V" in valeurs
    # Aucun prix dans le document.
    assert not any(isinstance(v, str) and "€" in v for v in valeurs)


def test_pdf_de_la_demande(client_essai: TestClient) -> None:
    numero = _demande(client_essai)
    reponse = client_essai.get(f"/api/commandes/{numero}/demande-devis.pdf")
    assert reponse.status_code == 200
    assert reponse.headers["content-type"] == "application/pdf"
    assert reponse.content.startswith(b"%PDF")


def test_ligne_annulee_absente_et_texte_hors_latin1(client_essai: TestClient) -> None:
    numero = _demande(client_essai)
    lignes = client_essai.get(f"/api/commandes/{numero}/lignes").json()
    client_essai.patch(
        f"/api/commandes/{numero}/lignes/{lignes[0]['id']}", json={"statut_ligne": "Annulee"}
    )
    client_essai.patch("/api/composants/ESSAI-AGI-001", json={"designation": "Cœur – « essai »"})
    conn = db.connect(client_essai.app.state.chemin_base)
    donnees = demande_document.get_donnees(conn, numero, {"nom": "A", "mail": None}, None)
    conn.close()
    assert donnees["lignes"][0]["designation"] == "Cœur – « essai »"
    assert len(donnees["lignes"]) == 2
    # Les caractères hors latin-1 sont remplacés au lieu de faire échouer le PDF.
    assert demande_document.build_pdf(donnees).startswith(b"%PDF")


def test_demandeur_connecte(client_essai: TestClient) -> None:
    conn = db.connect(client_essai.app.state.chemin_base)
    personne = Utilisateur(
        id=3, identifiant="camille@ecole.test", nom="Camille", role="contributeur"
    )
    assert demande_document.demandeur(conn, personne) == {
        "nom": "Camille",
        "mail": "camille@ecole.test",
    }
    conn.close()
