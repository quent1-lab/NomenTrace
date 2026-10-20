"""Partir de l'Excel d'une équipe : modèle vierge tous blocs et consigne pour un outil d'IA."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_import import _appliquer, _decisions_par_defaut, _deposer, _lignes


def _vierge(client: TestClient) -> bytes:
    reponse = client.get("/api/modeles/vierge")
    assert reponse.status_code == 200, reponse.text
    return reponse.content


def test_modele_vierge_sans_composant(client_essai: TestClient) -> None:
    classeur = load_workbook(io.BytesIO(_vierge(client_essai)))
    feuille = classeur["Composants"]
    entetes = [c.value for c in feuille[1]]
    assert entetes[:2] == ["ID", "Bloc"] and "Ensemble" in entetes
    assert all(c.value is None for ligne in feuille.iter_rows(min_row=2) for c in ligne)
    consignes = " ".join(str(c.value) for (c,) in classeur["Consignes"].iter_rows(max_col=1))
    assert "ainsi que le Bloc" in consignes


def test_modele_vierge_rempli_puis_importe(client_essai: TestClient) -> None:
    """Deux blocs dans un même fichier ; une ligne sans bloc est refusée."""
    classeur = load_workbook(io.BytesIO(_vierge(client_essai)))
    feuille = classeur["Composants"]
    entetes = [c.value for c in feuille[1]]

    def ligne(rang: int, **valeurs: object) -> None:
        for entete, valeur in valeurs.items():
            feuille.cell(rang, entetes.index(entete) + 1, valeur)

    ligne(
        2,
        **{
            "Bloc": "ALI",
            "Fonction": "Essai",
            "Désignation": "Câble secteur",
            "Mode appro": "Achat",
            "Qté besoin": 2,
        },
    )
    ligne(
        3,
        **{
            "Bloc": "OBS",
            "Fonction": "Essai",
            "Désignation": "Capteur de fin de course",
            "Mode appro": "Achat",
            "Qté besoin": 4,
        },
    )
    ligne(
        4,
        **{"Fonction": "Essai", "Désignation": "Sans bloc", "Mode appro": "Achat", "Qté besoin": 1},
    )
    flux = io.BytesIO()
    classeur.save(flux)
    depot = _deposer(client_essai, [("equipe.xlsx", flux.getvalue())])
    nouveaux = _lignes(depot, "NOUVEAU")
    assert {n["donnees"]["valeurs"]["bloc_code"] for n in nouveaux} == {"ALI", "OBS"}
    invalides = _lignes(depot, "INVALIDE")
    assert any("Bloc" in " ".join(i["donnees"]["raisons"]) for i in invalides)
    resultat = _appliquer(client_essai, depot, _decisions_par_defaut(depot))
    assert len(resultat["crees"]) == 2


def test_consigne_reprend_le_vocabulaire_du_projet(client_essai: TestClient) -> None:
    client_essai.post("/api/attributs", json={"libelle": "Tension", "type": "nombre", "unite": "V"})
    reponse = client_essai.get("/api/modeles/consigne-ia")
    assert reponse.status_code == 200
    consigne = reponse.json()["consigne"]
    blocs = client_essai.get("/api/blocs").json()
    assert all(f"- {b['code']} : {b['nom']}" in consigne for b in blocs)
    assert "Mode appro : " in consigne and "Achat" in consigne
    assert "Tension (V) : un nombre en V" in consigne
    assert "Laisse la colonne ID vide" in consigne
    # Le dépôt est public : la consigne ne nomme aucun outil.
    assert "Claude" not in consigne and "ChatGPT" not in consigne
